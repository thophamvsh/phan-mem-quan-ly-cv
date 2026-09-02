from io import BytesIO

from django.http import HttpResponse
from openpyxl import Workbook, load_workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter
from tablib import Dataset


VIETNAMESE_COLUMNS = (
    ("ma_nhan_vien", "Mã nhân viên", 18),
    ("ho_ten", "Họ và tên", 28),
    ("ma_don_vi", "Mã đơn vị", 18),
    ("ma_bo_phan", "Mã bộ phận", 18),
    ("username", "Tài khoản", 22),
    ("chuc_danh", "Chức danh", 22),
    ("dien_thoai", "Điện thoại", 16),
    ("tu_ngay", "Từ ngày", 14),
    ("den_ngay", "Đến ngày", 14),
    ("dang_lam_viec", "Đang làm việc", 16),
)
LABEL_TO_FIELD = {label: field for field, label, _ in VIETNAMESE_COLUMNS}


def _format_date(value):
    return value.strftime("%d/%m/%Y") if value else ""


def build_staff_workbook(queryset, units, departments, *, is_template=False):
    unit_rows = list(units)
    department_rows = list(departments)
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "Danh mục nhân sự"
    last_column = len(VIETNAMESE_COLUMNS)

    plant_names = sorted({
        unit.nha_may.ten_nha_may
        for unit in unit_rows
        if unit.nha_may_id
    })
    if len(plant_names) == 1:
        plant_name = plant_names[0].upper()
        plant_display = (
            plant_name
            if plant_name.startswith("NHÀ MÁY")
            else f"NHÀ MÁY THỦY ĐIỆN {plant_name}"
        )
    else:
        plant_display = "PHẠM VI CÁC NHÀ MÁY"

    left_end = 5
    sheet.merge_cells(start_row=1, start_column=1, end_row=1, end_column=left_end)
    company = sheet.cell(1, 1, "CÔNG TY CỔ PHẦN THỦY ĐIỆN VĨNH SƠN - SÔNG HINH")
    company.font = Font(name="Arial", size=10, bold=True, color="1E3A8A")
    company.alignment = Alignment(horizontal="center", vertical="center")
    sheet.merge_cells(start_row=2, start_column=1, end_row=2, end_column=left_end)
    plant = sheet.cell(2, 1, plant_display)
    plant.font = Font(name="Arial", size=10, bold=True, color="334155")
    plant.alignment = Alignment(horizontal="center", vertical="center")
    sheet.merge_cells(start_row=3, start_column=1, end_row=3, end_column=left_end)
    department_label = sheet.cell(3, 1, "Danh mục tổ chức và nhân sự")
    department_label.font = Font(name="Arial", size=9, italic=True, color="64748B")
    department_label.alignment = Alignment(horizontal="center", vertical="center")

    sheet.merge_cells(start_row=1, start_column=6, end_row=1, end_column=last_column)
    country = sheet.cell(1, 6, "CỘNG HÒA XÃ HỘI CHỦ NGHĨA VIỆT NAM")
    country.font = Font(name="Arial", size=10, bold=True, color="0F172A")
    country.alignment = Alignment(horizontal="center", vertical="center")
    sheet.merge_cells(start_row=2, start_column=6, end_row=2, end_column=last_column)
    motto = sheet.cell(2, 6, "Độc lập - Tự do - Hạnh phúc")
    motto.font = Font(name="Arial", size=10, bold=True, color="334155")
    motto.alignment = Alignment(horizontal="center", vertical="center")
    sheet.merge_cells(start_row=3, start_column=6, end_row=3, end_column=last_column)
    line = sheet.cell(3, 6, "—————————")
    line.font = Font(name="Arial", size=9, color="94A3B8")
    line.alignment = Alignment(horizontal="center", vertical="center")

    sheet.merge_cells(start_row=5, start_column=1, end_row=5, end_column=last_column)
    title = sheet.cell(5, 1, "MẪU NHẬP DANH MỤC NHÂN SỰ" if is_template else "DANH MỤC NHÂN SỰ")
    title.font = Font(name="Arial", size=14, bold=True, color="1E3A8A")
    title.alignment = Alignment(horizontal="center", vertical="center")
    sheet.row_dimensions[5].height = 28

    sheet.merge_cells(start_row=6, start_column=1, end_row=6, end_column=last_column)
    note = sheet.cell(
        6,
        1,
        "Mã nhân viên là bắt buộc và dùng để cập nhật dữ liệu. Ngày nhập dạng DD/MM/YYYY. "
        "Mã bộ phận phải thuộc đúng mã đơn vị.",
    )
    note.font = Font(name="Arial", italic=True, color="475569")
    note.fill = PatternFill("solid", fgColor="EFF6FF")
    note.alignment = Alignment(wrap_text=True, vertical="center")
    sheet.row_dimensions[6].height = 32

    header_row = 7
    thin = Side(style="thin", color="CBD5E1")
    for column, (_, label, width) in enumerate(VIETNAMESE_COLUMNS, 1):
        cell = sheet.cell(header_row, column, label)
        cell.font = Font(name="Arial", bold=True, color="FFFFFF")
        cell.fill = PatternFill("solid", fgColor="1D4ED8")
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        cell.border = Border(left=thin, right=thin, top=thin, bottom=thin)
        sheet.column_dimensions[get_column_letter(column)].width = width
    sheet.row_dimensions[header_row].height = 28

    for row_number, person in enumerate(queryset, header_row + 1):
        values = (
            person.ma_nhan_vien or "",
            person.ho_ten,
            person.don_vi.ma_don_vi,
            person.bo_phan.ma_bo_phan,
            person.user.username if person.user else "",
            person.chuc_danh,
            person.dien_thoai,
            _format_date(person.tu_ngay),
            _format_date(person.den_ngay),
            "Có" if person.dang_lam_viec else "Không",
        )
        for column, value in enumerate(values, 1):
            cell = sheet.cell(row_number, column, value)
            cell.font = Font(name="Arial", size=10, color="1E293B")
            cell.fill = PatternFill("solid", fgColor="FFFFFF" if row_number % 2 else "F8FAFC")
            cell.border = Border(left=thin, right=thin, top=thin, bottom=thin)
            cell.alignment = Alignment(vertical="center", wrap_text=True)

    sheet.freeze_panes = "A8"
    sheet.auto_filter.ref = f"A7:{get_column_letter(last_column)}{max(header_row + 1, sheet.max_row)}"
    sheet.sheet_view.showGridLines = False

    directory = workbook.create_sheet("Danh mục mã")
    directory.append(["Mã đơn vị", "Tên đơn vị", "Mã bộ phận", "Tên bộ phận"])
    directory.freeze_panes = "A2"
    for cell in directory[1]:
        cell.font = Font(name="Arial", bold=True, color="FFFFFF")
        cell.fill = PatternFill("solid", fgColor="475569")
    total = max(len(unit_rows), len(department_rows), 1)
    for index in range(total):
        unit = unit_rows[index] if index < len(unit_rows) else None
        department = department_rows[index] if index < len(department_rows) else None
        directory.append([
            unit.ma_don_vi if unit else "",
            unit.ten_don_vi if unit else "",
            department.ma_bo_phan if department else "",
            department.ten_bo_phan if department else "",
        ])
    for column, width in zip("ABCD", (18, 32, 18, 32)):
        directory.column_dimensions[column].width = width

    output = BytesIO()
    workbook.save(output)
    output.seek(0)
    return output.getvalue()


def workbook_response(content, filename):
    response = HttpResponse(
        content,
        content_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
    )
    response["Content-Disposition"] = f'attachment; filename="{filename}"'
    return response


def staff_dataset_from_upload(uploaded_file):
    workbook = load_workbook(uploaded_file, data_only=True, read_only=True)
    sheet = workbook["Danh mục nhân sự"] if "Danh mục nhân sự" in workbook.sheetnames else workbook.active
    header_row = None
    headers = None
    for row_number, values in enumerate(sheet.iter_rows(values_only=True), 1):
        normalized = [str(value).strip() if value is not None else "" for value in values]
        if "Mã nhân viên" in normalized and "Họ và tên" in normalized:
            header_row = row_number
            headers = normalized
            break
    if not header_row:
        raise ValueError("Không tìm thấy dòng tiêu đề Mã nhân viên/Họ và tên trong file Excel.")

    internal_headers = [LABEL_TO_FIELD.get(value, value) for value in headers]
    dataset = Dataset(headers=internal_headers)
    for values in sheet.iter_rows(min_row=header_row + 1, values_only=True):
        if not any(value not in (None, "") for value in values):
            continue
        row = list(values)
        active_index = internal_headers.index("dang_lam_viec")
        active_value = str(row[active_index] or "").strip().lower()
        if active_value in {"có", "co", "x", "yes", "true", "1"}:
            row[active_index] = "1"
        elif active_value in {"không", "khong", "no", "false", "0"}:
            row[active_index] = "0"
        dataset.append(row)
    return dataset
