from django.contrib.auth import get_user_model
from io import BytesIO
from openpyxl import load_workbook
from rest_framework import status
from rest_framework.test import APITestCase

from core.models import UserProfile
from tochuc.models import BoPhan, DonViToChuc, NhaMay, NhanSu


User = get_user_model()


def create_user(username, plant, **permissions):
    user = User.objects.create_user(
        username=username,
        email=f"{username}@example.com",
        password="test-pass-123",
    )
    profile, _ = UserProfile.objects.get_or_create(user=user)
    profile.nha_may = plant
    for field, value in permissions.items():
        setattr(profile, field, value)
    profile.save()
    return user


class OrganizationDirectoryAPITests(APITestCase):
    def setUp(self):
        self.song_hinh = NhaMay.objects.create(
            ma_nha_may="SH-API",
            ten_nha_may="Sông Hinh",
        )
        self.vinh_son = NhaMay.objects.create(
            ma_nha_may="VS-API",
            ten_nha_may="Vĩnh Sơn",
        )
        self.song_hinh_unit = DonViToChuc.objects.create(
            ma_don_vi="SH-UNIT",
            ten_don_vi="Quản lý vận hành SH",
            nha_may=self.song_hinh,
        )
        self.vinh_son_unit = DonViToChuc.objects.create(
            ma_don_vi="VS-UNIT",
            ten_don_vi="Quản lý vận hành VS",
            nha_may=self.vinh_son,
        )
        self.song_hinh_department = BoPhan.objects.create(
            don_vi=self.song_hinh_unit,
            ma_bo_phan="VH",
            ten_bo_phan="Vận hành",
        )

    def test_viewer_only_sees_own_factory(self):
        user = create_user(
            "org-viewer",
            self.song_hinh,
            can_view_organization_directory=True,
        )
        self.client.force_authenticate(user)

        response = self.client.get("/api/v1/tochuc/don-vi/")

        self.assertEqual(response.status_code, status.HTTP_200_OK)
        self.assertEqual(
            {item["id"] for item in response.data},
            {self.song_hinh_unit.id},
        )

    def test_factory_endpoint_is_authenticated_and_scoped(self):
        anonymous = self.client.get("/api/v1/tochuc/nha-may/")
        user = create_user(
            "factory-viewer",
            self.song_hinh,
            can_view_organization_directory=True,
        )
        self.client.force_authenticate(user)
        scoped = self.client.get("/api/v1/tochuc/nha-may/")

        self.assertEqual(anonymous.status_code, status.HTTP_401_UNAUTHORIZED)
        self.assertEqual(scoped.status_code, status.HTTP_200_OK)
        self.assertEqual(
            {item["id"] for item in scoped.data},
            {self.song_hinh.id},
        )
        self.assertNotIn("Deprecation", scoped)

    def test_legacy_factory_endpoint_points_to_shared_successor(self):
        response = self.client.get("/api/v1/khovattu/auth/nha-may/")

        self.assertEqual(response.status_code, status.HTTP_200_OK)
        self.assertEqual(response["Deprecation"], "true")
        self.assertEqual(
            response["Link"],
            '</api/v1/tochuc/nha-may/>; rel="successor-version"',
        )

    def test_user_without_related_permission_is_denied(self):
        user = create_user("org-denied", self.song_hinh)
        self.client.force_authenticate(user)

        response = self.client.get("/api/v1/tochuc/don-vi/")

        self.assertEqual(response.status_code, status.HTTP_403_FORBIDDEN)

    def test_manager_can_create_in_own_factory_but_not_another_factory(self):
        user = create_user(
            "org-manager",
            self.song_hinh,
            can_manage_organization_directory=True,
        )
        self.client.force_authenticate(user)

        own_response = self.client.post(
            "/api/v1/tochuc/don-vi/",
            {
                "ma_don_vi": "SH-NEW",
                "ten_don_vi": "Đơn vị mới SH",
                "nha_may": self.song_hinh.id,
            },
            format="json",
        )
        other_response = self.client.post(
            "/api/v1/tochuc/don-vi/",
            {
                "ma_don_vi": "VS-NEW",
                "ten_don_vi": "Đơn vị mới VS",
                "nha_may": self.vinh_son.id,
            },
            format="json",
        )

        self.assertEqual(own_response.status_code, status.HTTP_201_CREATED)
        self.assertEqual(other_response.status_code, status.HTTP_403_FORBIDDEN)

    def test_delete_marks_department_inactive_and_options_hide_it(self):
        user = create_user(
            "org-soft-delete",
            self.song_hinh,
            can_manage_organization_directory=True,
        )
        self.client.force_authenticate(user)

        delete_response = self.client.delete(
            f"/api/v1/tochuc/bo-phan/{self.song_hinh_department.id}/"
        )
        options_response = self.client.get(
            f"/api/v1/tochuc/bo-phan/options/?don_vi={self.song_hinh_unit.id}"
        )
        self.song_hinh_department.refresh_from_db()

        self.assertEqual(delete_response.status_code, status.HTTP_204_NO_CONTENT)
        self.assertFalse(self.song_hinh_department.dang_hoat_dong)
        self.assertEqual(options_response.status_code, status.HTTP_200_OK)
        self.assertEqual(options_response.data, [])

    def test_legacy_shift_endpoint_uses_same_shared_records(self):
        user = create_user(
            "legacy-org-viewer",
            self.song_hinh,
            can_view_shift_schedule=True,
        )
        self.client.force_authenticate(user)

        shared_response = self.client.get("/api/v1/tochuc/don-vi/")
        legacy_response = self.client.get("/api/v1/quanlycatruc/don-vi/")

        self.assertEqual(shared_response.status_code, status.HTTP_200_OK)
        self.assertEqual(legacy_response.status_code, status.HTTP_200_OK)
        self.assertEqual(legacy_response["Deprecation"], "true")
        self.assertEqual(
            legacy_response["Link"],
            '</api/v1/tochuc/don-vi/>; rel="successor-version"',
        )
        self.assertEqual(
            {item["id"] for item in shared_response.data},
            {item["id"] for item in legacy_response.data},
        )

    def test_shared_staff_endpoint_preserves_legacy_compatibility(self):
        user = create_user(
            "shared-staff-manager",
            self.song_hinh,
            can_manage_organization_directory=True,
        )
        self.client.force_authenticate(user)
        create_response = self.client.post(
            "/api/v1/tochuc/nhan-su/",
            {
                "ma_nhan_vien": "SH-NV-01",
                "ho_ten": "Nguyễn Văn Vận Hành",
                "don_vi": self.song_hinh_unit.id,
                "bo_phan": self.song_hinh_department.id,
                "chuc_danh": "Vận hành viên",
                "dang_lam_viec": True,
            },
            format="json",
        )
        shared_response = self.client.get("/api/v1/tochuc/nhan-su/")
        legacy_response = self.client.get(
            "/api/v1/quanlycatruc/nhan-su/"
        )
        account_options_response = self.client.get(
            "/api/v1/tochuc/nhan-su/tai-khoan-options/",
            {"nha_may": self.song_hinh.id},
        )

        self.assertEqual(create_response.status_code, status.HTTP_201_CREATED)
        self.assertIsNone(create_response.data["user"])
        self.assertEqual(shared_response.status_code, status.HTTP_200_OK)
        self.assertEqual(legacy_response.status_code, status.HTTP_200_OK)
        self.assertEqual(legacy_response["Deprecation"], "true")
        self.assertEqual(
            legacy_response["Link"],
            '</api/v1/tochuc/nhan-su/>; rel="successor-version"',
        )
        self.assertEqual(
            account_options_response.status_code,
            status.HTTP_200_OK,
        )
        self.assertEqual(
            {item["id"] for item in account_options_response.data},
            {user.id},
        )
        self.assertEqual(
            {item["id"] for item in shared_response.data},
            {item["id"] for item in legacy_response.data},
        )
        self.assertEqual(NhanSu.objects.count(), 1)

    def test_shared_staff_is_scoped_by_factory(self):
        vinh_son_department = BoPhan.objects.create(
            don_vi=self.vinh_son_unit,
            ma_bo_phan="VH",
            ten_bo_phan="Vận hành",
        )
        own_staff = NhanSu.objects.create(
            ho_ten="Nhân sự Sông Hinh",
            don_vi=self.song_hinh_unit,
            bo_phan=self.song_hinh_department,
        )
        NhanSu.objects.create(
            ho_ten="Nhân sự Vĩnh Sơn",
            don_vi=self.vinh_son_unit,
            bo_phan=vinh_son_department,
        )
        user = create_user(
            "shared-staff-viewer",
            self.song_hinh,
            can_view_organization_directory=True,
        )
        self.client.force_authenticate(user)

        response = self.client.get("/api/v1/tochuc/nhan-su/")

        self.assertEqual(response.status_code, status.HTTP_200_OK)
        self.assertEqual(
            {item["id"] for item in response.data},
            {own_staff.id},
        )

    def test_staff_excel_template_and_export_have_vietnamese_headers(self):
        user = create_user(
            "staff-excel-viewer",
            self.song_hinh,
            can_view_organization_directory=True,
        )
        NhanSu.objects.create(
            ma_nhan_vien="SH-XLSX-01",
            ho_ten="Nguyễn Văn Excel",
            don_vi=self.song_hinh_unit,
            bo_phan=self.song_hinh_department,
        )
        self.client.force_authenticate(user)

        template = self.client.get("/api/v1/tochuc/nhan-su/excel-template/")
        exported = self.client.get("/api/v1/tochuc/nhan-su/export-excel/")

        self.assertEqual(template.status_code, status.HTTP_200_OK)
        self.assertEqual(exported.status_code, status.HTTP_200_OK)
        workbook = load_workbook(BytesIO(exported.content), data_only=True)
        sheet = workbook["Danh mục nhân sự"]
        self.assertEqual(sheet["A1"].value, "CÔNG TY CỔ PHẦN THỦY ĐIỆN VĨNH SƠN - SÔNG HINH")
        self.assertEqual(sheet["F1"].value, "CỘNG HÒA XÃ HỘI CHỦ NGHĨA VIỆT NAM")
        self.assertEqual(sheet["A7"].value, "Mã nhân viên")
        self.assertEqual(sheet["B7"].value, "Họ và tên")
        self.assertEqual(sheet["C7"].value, "Mã nhà máy")
        self.assertEqual(sheet["C8"].value, self.song_hinh.ma_nha_may)
        self.assertEqual(sheet["B8"].value, "Nguyễn Văn Excel")

    def test_manager_imports_staff_excel_in_own_factory(self):
        user = create_user(
            "staff-excel-manager",
            self.song_hinh,
            can_manage_organization_directory=True,
        )
        self.client.force_authenticate(user)
        template = self.client.get("/api/v1/tochuc/nhan-su/excel-template/")
        workbook = load_workbook(BytesIO(template.content))
        sheet = workbook["Danh mục nhân sự"]
        sheet.append([
            "SH-XLSX-NEW",
            "Trần Văn Nhập",
            self.song_hinh.ma_nha_may,
            self.song_hinh.ten_nha_may,
            self.song_hinh_unit.ma_don_vi,
            self.song_hinh_department.ma_bo_phan,
            "",
            "Trực chính",
            "0900000000",
            "01/09/2026",
            "",
            "Có",
        ])
        payload = BytesIO()
        workbook.save(payload)
        payload.seek(0)
        payload.name = "nhan-su.xlsx"

        response = self.client.post(
            "/api/v1/tochuc/nhan-su/import-excel/",
            {"file": payload},
            format="multipart",
        )

        self.assertEqual(response.status_code, status.HTTP_200_OK, response.data)
        self.assertEqual(response.data["created"], 1)
        self.assertTrue(NhanSu.objects.filter(ma_nhan_vien="SH-XLSX-NEW").exists())

    def test_staff_excel_import_rejects_unit_outside_factory(self):
        user = create_user(
            "staff-excel-scoped",
            self.song_hinh,
            can_manage_organization_directory=True,
        )
        self.client.force_authenticate(user)
        workbook = load_workbook(BytesIO(
            self.client.get("/api/v1/tochuc/nhan-su/excel-template/").content
        ))
        sheet = workbook["Danh mục nhân sự"]
        sheet.append([
            "VS-XLSX-BLOCKED",
            "Nhân sự ngoài phạm vi",
            self.vinh_son.ma_nha_may,
            self.vinh_son.ten_nha_may,
            self.vinh_son_unit.ma_don_vi,
            "VH",
            "",
            "",
            "",
            "01/09/2026",
            "",
            "Có",
        ])
        payload = BytesIO()
        workbook.save(payload)
        payload.seek(0)
        payload.name = "nhan-su.xlsx"

        response = self.client.post(
            "/api/v1/tochuc/nhan-su/import-excel/",
            {"file": payload},
            format="multipart",
        )

        self.assertEqual(response.status_code, status.HTTP_400_BAD_REQUEST)
        self.assertFalse(NhanSu.objects.filter(ma_nhan_vien="VS-XLSX-BLOCKED").exists())
