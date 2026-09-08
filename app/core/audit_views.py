import io
from auditlog.models import LogEntry
from django.contrib.contenttypes.models import ContentType
from django.db.models import Q
from django.http import HttpResponse
from django.utils import timezone
from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from rest_framework import permissions, status, viewsets
from rest_framework.exceptions import PermissionDenied, ValidationError
from rest_framework.pagination import PageNumberPagination
from rest_framework.response import Response
from rest_framework.views import APIView

from tochuc.models import NhaMay
from .models import UserActivityLog, UserManagementAudit
from .throttles import AuditExportRateThrottle
from .audit_serializers import (
    ACTION_DISPLAY_MAP,
    AuditExportRequestSerializer,
    AuditFilterSerializer,
    DataChangeLogSerializer,
    MODEL_DISPLAY_NAMES,
    USER_MGT_ACTION_MAP,
    UserActivityLogSerializer,
    UserManagementAuditSerializer,
    format_user_display,
    latest_change_value,
    sanitize_changes,
    sanitize_text,
)


ACTIVITY_ACTIONS = {'LOGIN', 'LOGOUT', 'LOGIN_FAILED'}
USER_MANAGEMENT_ACTIONS = {'create', 'edit', 'role', 'status'}
DATA_ACTIONS = {0, 1, 2, 3}


def has_profile_permission(user, permission_name):
    return bool(
        user
        and user.is_authenticated
        and (
            user.is_superuser
            or (
                hasattr(user, 'profile')
                and user.profile.has_effective_permission(permission_name)
            )
        )
    )


def validate_filters(raw_data, tab=None, export=False):
    # Các bộ lọc là tùy chọn. Frontend cũ có thể gửi chuỗi rỗng cho select
    # số; loại bỏ trước khi DRF IntegerField kiểm tra để giữ tương thích.
    cleaned_data = raw_data.copy()
    for field_name in (
        'plant_id', 'nha_may', 'action', 'action_type', 'model',
        'model_name', 'content_type_id', 'from_date', 'to_date', 'search',
    ):
        if cleaned_data.get(field_name) == '':
            cleaned_data.pop(field_name, None)
    serializer_class = AuditExportRequestSerializer if export else AuditFilterSerializer
    serializer = serializer_class(data=cleaned_data)
    serializer.is_valid(raise_exception=True)
    filters = serializer.validated_data
    tab = filters.get('tab', tab)
    action = filters.get('action') or filters.get('action_type')
    if action not in (None, ''):
        if tab == 'data':
            try:
                parsed_action = int(action)
            except (TypeError, ValueError) as exc:
                raise ValidationError({'action': 'Hành động dữ liệu không hợp lệ.'}) from exc
            if parsed_action not in DATA_ACTIONS:
                raise ValidationError({'action': 'Hành động dữ liệu không hợp lệ.'})
            filters['action'] = parsed_action
        else:
            allowed = ACTIVITY_ACTIONS if tab == 'activity' else USER_MANAGEMENT_ACTIONS
            normalized = str(action).upper() if tab == 'activity' else str(action).lower()
            if normalized not in allowed:
                raise ValidationError({'action': 'Hành động kiểm toán không hợp lệ.'})
            filters['action'] = normalized
    return filters


def get_scope(user):
    is_super = user.is_superuser
    is_all = is_super or (
        hasattr(user, 'profile') and user.profile.is_all_factories
    )
    plant_id = (
        getattr(user.profile, 'nha_may_id', None)
        if hasattr(user, 'profile') else None
    )
    return is_super, is_all, plant_id


class AuditLogPagination(PageNumberPagination):
    page_size = 25
    page_size_query_param = 'page_size'
    max_page_size = 100


class HasActivityLogPermission(permissions.BasePermission):
    def has_permission(self, request, view):
        return has_profile_permission(request.user, 'can_view_activity_logs')


class HasDataAuditLogPermission(permissions.BasePermission):
    def has_permission(self, request, view):
        return has_profile_permission(request.user, 'can_view_data_audit_logs')


class HasUserManagementAuditPermission(permissions.BasePermission):
    def has_permission(self, request, view):
        return has_profile_permission(request.user, 'can_view_user_management_audit')


class HasExportAuditLogPermission(permissions.BasePermission):
    def has_permission(self, request, view):
        return has_profile_permission(request.user, 'can_export_audit_logs')


class HasAnyAuditViewPermission(permissions.BasePermission):
    def has_permission(self, request, view):
        return any(
            has_profile_permission(request.user, permission_name)
            for permission_name in (
                'can_view_activity_logs',
                'can_view_data_audit_logs',
                'can_view_user_management_audit',
            )
        )


class UserActivityLogViewSet(viewsets.ReadOnlyModelViewSet):
    """
    API chỉ đọc tra cứu nhật ký đăng nhập và truy cập mạng (UserActivityLog).
    Bảo vệ bằng quyền can_view_activity_logs và phân vùng nhà máy nghiêm ngặt.
    """
    serializer_class = UserActivityLogSerializer
    pagination_class = AuditLogPagination
    permission_classes = [permissions.IsAuthenticated, HasActivityLogPermission]

    def get_queryset(self):
        user = self.request.user
        filters = validate_filters(self.request.query_params, tab='activity')
        qs = UserActivityLog.objects.select_related('user', 'nha_may').order_by('-timestamp')

        # 1. Phân vùng nhà máy
        is_super, is_all, user_plant = get_scope(user)
        if not is_all:
            if user_plant:
                qs = qs.filter(nha_may_id=user_plant)
            else:
                # Không có nhà máy và không có quyền toàn nhà máy -> không thấy log
                return qs.none()
        else:
            if not is_super:
                qs = qs.filter(nha_may_id__isnull=False)
            # Tài khoản toàn quyền có thể lọc theo nhà máy cụ thể nếu truyền query param
            plant_id = filters.get('plant_id') or filters.get('nha_may')
            if plant_id:
                qs = qs.filter(nha_may_id=plant_id)

        # 2. Lọc theo loại hành động
        action = filters.get('action')
        if action:
            qs = qs.filter(action_type__iexact=action)

        # 3. Lọc theo khoảng thời gian
        from_date = filters.get('from_date')
        if from_date:
            qs = qs.filter(timestamp__date__gte=from_date)

        to_date = filters.get('to_date')
        if to_date:
            qs = qs.filter(timestamp__date__lte=to_date)

        # 4. Tìm kiếm từ khóa (username, email, description, IP)
        search = filters.get('search')
        if search:
            search = search.strip()
            qs = qs.filter(
                Q(user__username__icontains=search) |
                Q(user__email__icontains=search) |
                Q(description__icontains=search) |
                Q(ip_address__icontains=search)
            )

        return qs


class DataChangeLogViewSet(viewsets.ReadOnlyModelViewSet):
    """
    API chỉ đọc tra cứu vết thay đổi dữ liệu nghiệp vụ (auditlog.LogEntry).
    Bảo vệ bằng quyền can_view_data_audit_logs và phân vùng dữ liệu qua plant snapshot.
    """
    serializer_class = DataChangeLogSerializer
    pagination_class = AuditLogPagination
    permission_classes = [permissions.IsAuthenticated, HasDataAuditLogPermission]

    def get_queryset(self):
        user = self.request.user
        filters = validate_filters(self.request.query_params, tab='data')
        qs = LogEntry.objects.select_related('content_type', 'actor').order_by('-timestamp')

        # 1. Phân vùng nhà máy dựa vào snapshot additional_data['plant_id']
        is_super, is_all, user_plant = get_scope(user)

        plant_param = filters.get('plant_id') or filters.get('nha_may')

        if not is_all:
            if not user_plant:
                return qs.none()
            # Quản lý 1 nhà máy: CHỈ xem các log có additional_data.plant_id = user_plant
            qs = qs.filter(additional_data__plant_id=int(user_plant))
        else:
            if not is_super:
                # is_all_factories nhưng không phải superuser: xem được tất cả các nhà máy, nhưng KHÔNG xem log unscoped
                qs = qs.filter(additional_data__plant_id__isnull=False)

            if plant_param:
                qs = qs.filter(additional_data__plant_id=int(plant_param))

        # 2. Lọc theo model / phân hệ
        model_name = filters.get('model') or filters.get('model_name')
        if model_name:
            qs = qs.filter(content_type__model__iexact=model_name)

        content_type_id = filters.get('content_type_id')
        if content_type_id:
            qs = qs.filter(content_type_id=content_type_id)

        # 3. Lọc theo thao tác (0: Tạo, 1: Sửa, 2: Xóa)
        action = filters.get('action')
        if action is not None:
            qs = qs.filter(action=action)

        # 4. Lọc theo khoảng ngày
        from_date = filters.get('from_date')
        if from_date:
            qs = qs.filter(timestamp__date__gte=from_date)

        to_date = filters.get('to_date')
        if to_date:
            qs = qs.filter(timestamp__date__lte=to_date)

        # 5. Tìm kiếm
        search = filters.get('search')
        if search:
            search = search.strip()
            qs = qs.filter(
                Q(object_repr__icontains=search) |
                Q(actor__username__icontains=search) |
                Q(actor__email__icontains=search)
            )

        return qs


class UserManagementAuditViewSet(viewsets.ReadOnlyModelViewSet):
    """
    API chỉ đọc tra cứu lịch sử quản lý tài khoản (UserManagementAudit).
    Bảo vệ bằng quyền can_view_user_management_audit và phân vùng theo nhà máy của target.
    """
    serializer_class = UserManagementAuditSerializer
    pagination_class = AuditLogPagination
    permission_classes = [permissions.IsAuthenticated, HasUserManagementAuditPermission]

    def get_queryset(self):
        user = self.request.user
        filters = validate_filters(self.request.query_params, tab='user_management')
        qs = UserManagementAudit.objects.select_related('actor', 'target', 'nha_may').order_by('-created_at')

        # 1. Phân vùng nhà máy
        is_super, is_all, user_plant = get_scope(user)
        if not is_all:
            if user_plant:
                qs = qs.filter(nha_may_id=user_plant)
            else:
                return qs.none()
        else:
            if not is_super:
                qs = qs.filter(nha_may_id__isnull=False)
            plant_id = filters.get('plant_id') or filters.get('nha_may')
            if plant_id:
                qs = qs.filter(nha_may_id=plant_id)

        # 2. Lọc theo hành động (create, edit, role, status)
        action = filters.get('action')
        if action:
            qs = qs.filter(action__iexact=action)

        # 3. Lọc theo khoảng ngày
        from_date = filters.get('from_date')
        if from_date:
            qs = qs.filter(created_at__date__gte=from_date)

        to_date = filters.get('to_date')
        if to_date:
            qs = qs.filter(created_at__date__lte=to_date)

        # 4. Tìm kiếm (username actor hoặc target)
        search = filters.get('search')
        if search:
            search = search.strip()
            qs = qs.filter(
                Q(actor__username__icontains=search) |
                Q(actor__email__icontains=search) |
                Q(target__username__icontains=search) |
                Q(target__email__icontains=search)
            )

        return qs


class AuditMetadataView(APIView):
    """
    Trả về danh mục cấu hình bộ lọc (nhà máy, danh sách model, danh sách thao tác)
    trong phạm vi thẩm quyền của người dùng hiện tại, không rò rỉ thông tin ra ngoài.
    """
    permission_classes = [permissions.IsAuthenticated, HasAnyAuditViewPermission]

    def get(self, request):
        user = request.user
        is_super, is_all, user_plant_id = get_scope(user)

        # 1. Danh sách nhà máy được phép
        if is_all:
            plants = list(NhaMay.objects.values('id', 'ma_nha_may', 'ten_nha_may').order_by('id'))
        elif user_plant_id:
            plants = list(NhaMay.objects.filter(id=user_plant_id).values('id', 'ma_nha_may', 'ten_nha_may'))
        else:
            plants = []

        # 2. Danh sách model có trong auditlog
        # Lấy các ContentType thực tế đang có log trong hệ thống
        available_models = []
        log_qs = LogEntry.objects.all()
        if not is_all:
            log_qs = (
                log_qs.filter(additional_data__plant_id=user_plant_id)
                if user_plant_id else log_qs.none()
            )
        elif not is_super:
            log_qs = log_qs.filter(additional_data__plant_id__isnull=False)
        if not has_profile_permission(user, 'can_view_data_audit_logs'):
            log_qs = log_qs.none()
        ct_ids = log_qs.values_list('content_type_id', flat=True).distinct()
        content_types = ContentType.objects.filter(id__in=ct_ids).order_by('model')

        for ct in content_types:
            model_key = ct.model.lower()
            display_name = MODEL_DISPLAY_NAMES.get(model_key, ct.name)
            available_models.append({
                'id': ct.id,
                'model': model_key,
                'app_label': ct.app_label,
                'name': display_name,
            })

        # Sắp xếp theo tên tiếng Việt
        available_models.sort(key=lambda x: x['name'])

        # 3. Danh sách loại hành động
        activity_actions = [
            {'code': 'LOGIN', 'name': 'Đăng nhập'},
            {'code': 'LOGOUT', 'name': 'Đăng xuất'},
            {'code': 'LOGIN_FAILED', 'name': 'Đăng nhập thất bại'},
        ]

        data_change_actions = [
            {'code': 0, 'name': 'Tạo mới'},
            {'code': 1, 'name': 'Cập nhật'},
            {'code': 2, 'name': 'Xóa'},
        ]

        user_management_actions = [
            {'code': 'create', 'name': 'Tạo tài khoản'},
            {'code': 'edit', 'name': 'Sửa thông tin'},
            {'code': 'role', 'name': 'Đổi vai trò'},
            {'code': 'status', 'name': 'Khóa/Mở tài khoản'},
        ]

        return Response({
            'plants': plants,
            'models': available_models,
            'activity_actions': activity_actions,
            'data_change_actions': data_change_actions,
            'user_management_actions': user_management_actions,
            'user_scope': {
                'is_superuser': is_super,
                'is_all_factories': is_all,
                'plant_id': user_plant_id,
            }
        })


def sanitize_excel_value(val):
    """
    Phòng chống tấn công tiêm công thức (CSV / Excel Formula Injection).
    Kiểm tra chuỗi: nếu bắt đầu bằng =, +, -, @, \t, \r (kể cả khi có khoảng trắng đầu chuỗi),
    tự động thêm tiền tố dấu nháy đơn ' ở đầu chuỗi để Excel coi là text thuần túy.
    """
    if val is None:
        return ""
    s = str(val)
    if s.startswith(('\t', '\r')):
        return "'" + s
    stripped = s.lstrip(' ')
    if stripped and stripped[0] in ('=', '+', '-', '@', '\t', '\r'):
        return "'" + s
    return s


class AuditExportExcelView(APIView):
    """
    API xuất dữ liệu báo cáo kiểm toán độc lập ra file Excel (.xlsx).
    - Bảo vệ nghiêm ngặt bằng quyền can_export_audit_logs.
    - Áp dụng AuditExportRateThrottle (tối đa 10 req/phút).
    - Khống chế an toàn tối đa 5.000 dòng dữ liệu.
    - Áp dụng cơ chế khử Formula Injection trên 100% các ô dữ liệu dạng chuỗi.
    """
    permission_classes = [permissions.IsAuthenticated, HasExportAuditLogPermission]
    throttle_classes = [AuditExportRateThrottle]

    def post(self, request):
        request_data = request.data.copy()
        if not request_data.get('tab'):
            request_data['tab'] = 'activity'
        tab = request_data.get('tab')
        filters = validate_filters(request_data, tab=tab, export=True)
        tab = filters['tab']
        user = request.user
        required_view_permission = {
            'activity': 'can_view_activity_logs',
            'data': 'can_view_data_audit_logs',
            'user_management': 'can_view_user_management_audit',
        }[tab]
        if not has_profile_permission(user, required_view_permission):
            raise PermissionDenied('Bạn không có quyền xem loại nhật ký cần xuất.')

        is_super, is_all, user_plant_id = get_scope(user)
        from_date = filters.get('from_date')
        to_date = filters.get('to_date')
        plant_param = filters.get('plant_id') or filters.get('nha_may')
        action_param = filters.get('action')
        search = filters.get('search', '')

        wb = Workbook()
        ws = wb.active

        # Style chung
        header_font = Font(name='Arial', size=11, bold=True, color='FFFFFF')
        header_fill = PatternFill(start_color='1E3A8A', end_color='1E3A8A', fill_type='solid')
        header_align = Alignment(horizontal='center', vertical='center', wrap_text=True)
        thin_border = Border(
            left=Side(style='thin', color='CBD5E1'),
            right=Side(style='thin', color='CBD5E1'),
            top=Side(style='thin', color='CBD5E1'),
            bottom=Side(style='thin', color='CBD5E1'),
        )
        content_font = Font(name='Arial', size=10)
        center_align = Alignment(horizontal='center', vertical='center')
        left_align = Alignment(horizontal='left', vertical='center')

        MAX_EXPORT_ROWS = 5000
        is_truncated = False
        total_records = 0

        if tab == 'activity':
            ws.title = "Nhật ký truy cập"
            headers = [
                "STT", "Thời gian", "Tên đăng nhập", "Họ tên",
                "Hành động", "Địa chỉ IP", "Thiết bị / Trình duyệt",
                "Nhà máy", "Mô tả"
            ]

            qs = UserActivityLog.objects.select_related('user', 'nha_may').order_by('-timestamp')
            if not is_all:
                if not user_plant_id:
                    qs = qs.none()
                else:
                    qs = qs.filter(nha_may_id=user_plant_id)
            else:
                if not is_super:
                    qs = qs.filter(nha_may_id__isnull=False)
                if plant_param:
                    qs = qs.filter(nha_may_id=plant_param)

            if action_param:
                qs = qs.filter(action_type__iexact=action_param)
            if from_date:
                qs = qs.filter(timestamp__date__gte=from_date)
            if to_date:
                qs = qs.filter(timestamp__date__lte=to_date)
            if search:
                qs = qs.filter(
                    Q(user__username__icontains=search) |
                    Q(user__email__icontains=search) |
                    Q(description__icontains=search) |
                    Q(ip_address__icontains=search)
                )

            total_records = qs.count()
            if total_records > MAX_EXPORT_ROWS:
                is_truncated = True
            records = list(qs[:MAX_EXPORT_ROWS])

            start_row = 1
            if is_truncated:
                ws.merge_cells(start_row=1, start_column=1, end_row=1, end_column=len(headers))
                warn_cell = ws.cell(row=1, column=1)
                warn_cell.value = (
                    f"CẢNH BÁO: Dữ liệu truy vấn có tổng cộng {total_records} dòng, vượt quá giới hạn 5.000 dòng. "
                    f"File này chỉ chứa 5.000 bản ghi đầu tiên. Vui lòng thu hẹp khoảng thời gian hoặc áp dụng thêm bộ lọc."
                )
                warn_cell.font = Font(name='Arial', size=10, bold=True, color='92400E')
                warn_cell.fill = PatternFill(start_color='FEF3C7', end_color='FEF3C7', fill_type='solid')
                warn_cell.alignment = Alignment(horizontal='center', vertical='center')
                start_row = 3

            for col_idx, h in enumerate(headers, 1):
                cell = ws.cell(row=start_row, column=col_idx, value=h)
                cell.font = header_font
                cell.fill = header_fill
                cell.alignment = header_align
                cell.border = thin_border
            ws.row_dimensions[start_row].height = 28

            for idx, item in enumerate(records, 1):
                r_idx = start_row + idx
                time_str = timezone.localtime(item.timestamp).strftime('%d/%m/%Y %H:%M:%S') if item.timestamp else ''
                action_display = {
                    'LOGIN': 'Đăng nhập',
                    'LOGOUT': 'Đăng xuất',
                    'LOGIN_FAILED': 'Đăng nhập thất bại'
                }.get(item.action_type, item.action_type)

                row_data = [
                    idx,
                    time_str,
                    sanitize_excel_value(item.user.username if item.user else "—"),
                    sanitize_excel_value(format_user_display(item.user)),
                    sanitize_excel_value(action_display),
                    sanitize_excel_value(item.ip_address or "—"),
                    sanitize_excel_value(sanitize_text(item.user_agent or "—")),
                    sanitize_excel_value(item.nha_may.ten_nha_may if item.nha_may else "Toàn cục"),
                    sanitize_excel_value(sanitize_text(item.description or "—")),
                ]
                for c_idx, val in enumerate(row_data, 1):
                    cell = ws.cell(row=r_idx, column=c_idx, value=val)
                    cell.font = content_font
                    cell.border = thin_border
                    cell.alignment = center_align if c_idx in (1, 2, 5, 6, 8) else left_align

        elif tab == 'data':
            ws.title = "Vết thay đổi dữ liệu"
            headers = [
                "STT", "Thời gian", "Người thực hiện", "Phân hệ / Model",
                "Bản ghi tác động", "Thao tác", "Nhà máy", "Chi tiết thay đổi"
            ]

            qs = LogEntry.objects.select_related('content_type', 'actor').order_by('-timestamp')
            if not is_all:
                if not user_plant_id:
                    qs = qs.none()
                else:
                    qs = qs.filter(additional_data__plant_id=int(user_plant_id))
            else:
                if not is_super:
                    qs = qs.filter(additional_data__plant_id__isnull=False)
                if plant_param:
                    qs = qs.filter(additional_data__plant_id=int(plant_param))

            model_name = filters.get('model') or filters.get('model_name')
            if model_name:
                qs = qs.filter(content_type__model__iexact=model_name)
            content_type_id = filters.get('content_type_id')
            if content_type_id:
                qs = qs.filter(content_type_id=content_type_id)
            if action_param is not None:
                qs = qs.filter(action=action_param)
            if from_date:
                qs = qs.filter(timestamp__date__gte=from_date)
            if to_date:
                qs = qs.filter(timestamp__date__lte=to_date)
            if search:
                qs = qs.filter(
                    Q(object_repr__icontains=search) |
                    Q(actor__username__icontains=search) |
                    Q(actor__email__icontains=search)
                )

            total_records = qs.count()
            if total_records > MAX_EXPORT_ROWS:
                is_truncated = True
            records = list(qs[:MAX_EXPORT_ROWS])

            start_row = 1
            if is_truncated:
                ws.merge_cells(start_row=1, start_column=1, end_row=1, end_column=len(headers))
                warn_cell = ws.cell(row=1, column=1)
                warn_cell.value = (
                    f"CẢNH BÁO: Dữ liệu truy vấn có tổng cộng {total_records} dòng, vượt quá giới hạn 5.000 dòng. "
                    f"File này chỉ chứa 5.000 bản ghi đầu tiên. Vui lòng thu hẹp khoảng thời gian hoặc chọn phân hệ cụ thể."
                )
                warn_cell.font = Font(name='Arial', size=10, bold=True, color='92400E')
                warn_cell.fill = PatternFill(start_color='FEF3C7', end_color='FEF3C7', fill_type='solid')
                warn_cell.alignment = Alignment(horizontal='center', vertical='center')
                start_row = 3

            for col_idx, h in enumerate(headers, 1):
                cell = ws.cell(row=start_row, column=col_idx, value=h)
                cell.font = header_font
                cell.fill = header_fill
                cell.alignment = header_align
                cell.border = thin_border
            ws.row_dimensions[start_row].height = 28

            for idx, item in enumerate(records, 1):
                r_idx = start_row + idx
                time_str = timezone.localtime(item.timestamp).strftime('%d/%m/%Y %H:%M:%S') if item.timestamp else ''
                model_key = item.content_type.model.lower() if item.content_type else ''
                model_display = MODEL_DISPLAY_NAMES.get(model_key, item.content_type.name if item.content_type else model_key)
                action_display = ACTION_DISPLAY_MAP.get(item.action, str(item.action))

                changes = sanitize_changes(item.changes)
                changes_text_list = []
                for f, diff in changes.items():
                    if isinstance(diff, (list, tuple)) and len(diff) == 2:
                        changes_text_list.append(f"{f}: {diff[0]} -> {diff[1]}")
                    else:
                        changes_text_list.append(f"{f}: {diff}")
                changes_str = "; ".join(changes_text_list) if changes_text_list else "—"

                add_data = item.additional_data or {}
                plant_code = add_data.get('plant_code') if isinstance(add_data, dict) else ''

                row_data = [
                    idx,
                    time_str,
                    sanitize_excel_value(format_user_display(item.actor)),
                    sanitize_excel_value(model_display),
                    sanitize_excel_value(sanitize_text(item.object_repr or "—")),
                    sanitize_excel_value(action_display),
                    sanitize_excel_value(plant_code or "Toàn cục"),
                    sanitize_excel_value(changes_str),
                ]
                for c_idx, val in enumerate(row_data, 1):
                    cell = ws.cell(row=r_idx, column=c_idx, value=val)
                    cell.font = content_font
                    cell.border = thin_border
                    cell.alignment = center_align if c_idx in (1, 2, 6, 7) else left_align

        elif tab == 'user_management':
            ws.title = "Quản trị tài khoản"
            headers = [
                "STT", "Thời gian", "Người quản trị", "Tài khoản mục tiêu",
                "Hành động", "Nhà máy", "Nội dung thay đổi"
            ]

            qs = UserManagementAudit.objects.select_related('actor', 'target', 'nha_may').order_by('-created_at')
            if not is_all:
                if not user_plant_id:
                    qs = qs.none()
                else:
                    qs = qs.filter(nha_may_id=user_plant_id)
            else:
                if not is_super:
                    qs = qs.filter(nha_may_id__isnull=False)
                if plant_param:
                    qs = qs.filter(nha_may_id=plant_param)

            if action_param:
                qs = qs.filter(action__iexact=action_param)
            if from_date:
                qs = qs.filter(created_at__date__gte=from_date)
            if to_date:
                qs = qs.filter(created_at__date__lte=to_date)
            if search:
                qs = qs.filter(
                    Q(actor__username__icontains=search) |
                    Q(actor__email__icontains=search) |
                    Q(target__username__icontains=search) |
                    Q(target__email__icontains=search)
                )

            total_records = qs.count()
            if total_records > MAX_EXPORT_ROWS:
                is_truncated = True
            records = list(qs[:MAX_EXPORT_ROWS])

            start_row = 1
            if is_truncated:
                ws.merge_cells(start_row=1, start_column=1, end_row=1, end_column=len(headers))
                warn_cell = ws.cell(row=1, column=1)
                warn_cell.value = (
                    f"CẢNH BÁO: Dữ liệu truy vấn có tổng cộng {total_records} dòng, vượt quá giới hạn 5.000 dòng. "
                    f"File này chỉ chứa 5.000 bản ghi đầu tiên. Vui lòng thu hẹp khoảng thời gian."
                )
                warn_cell.font = Font(name='Arial', size=10, bold=True, color='92400E')
                warn_cell.fill = PatternFill(start_color='FEF3C7', end_color='FEF3C7', fill_type='solid')
                warn_cell.alignment = Alignment(horizontal='center', vertical='center')
                start_row = 3

            for col_idx, h in enumerate(headers, 1):
                cell = ws.cell(row=start_row, column=col_idx, value=h)
                cell.font = header_font
                cell.fill = header_fill
                cell.alignment = header_align
                cell.border = thin_border
            ws.row_dimensions[start_row].height = 28

            for idx, item in enumerate(records, 1):
                r_idx = start_row + idx
                time_str = timezone.localtime(item.created_at).strftime('%d/%m/%Y %H:%M:%S') if item.created_at else ''
                action_display = USER_MGT_ACTION_MAP.get(item.action, item.action)

                changes = sanitize_changes(item.changes)
                if item.action == 'create':
                    role_id = latest_change_value(changes.get('role_id'))
                    changes_str = f"Tạo tài khoản mới (Vai trò ID: {role_id or 'Mặc định'})"
                elif item.action == 'status':
                    is_active = latest_change_value(changes.get('is_active'))
                    changes_str = f"Trạng thái: {'Hoạt động' if is_active else 'Bị khóa'}"
                elif item.action == 'role':
                    role_id = latest_change_value(changes.get('role_id'))
                    changes_str = f"Gán vai trò mới (ID vai trò: {role_id})"
                elif item.action == 'edit':
                    changes_str = f"Cập nhật: {', '.join(changes.keys())}" if changes else "Sửa thông tin"
                else:
                    changes_str = str(changes)

                row_data = [
                    idx,
                    time_str,
                    sanitize_excel_value(format_user_display(item.actor)),
                    sanitize_excel_value(format_user_display(item.target)),
                    sanitize_excel_value(action_display),
                    sanitize_excel_value(item.nha_may.ten_nha_may if item.nha_may else "Toàn cục"),
                    sanitize_excel_value(changes_str),
                ]
                for c_idx, val in enumerate(row_data, 1):
                    cell = ws.cell(row=r_idx, column=c_idx, value=val)
                    cell.font = content_font
                    cell.border = thin_border
                    cell.alignment = center_align if c_idx in (1, 2, 5, 6) else left_align

        else:
            return Response(
                {'detail': f'Loại báo cáo "{tab}" không được hỗ trợ.'},
                status=status.HTTP_400_BAD_REQUEST
            )

        # Căn chỉnh độ rộng cột tự động
        for col in ws.columns:
            # Lấy độ dài lớn nhất của text trong cột
            max_len = 0
            for cell in col:
                val_str = str(cell.value or '')
                if len(val_str) > max_len and cell.coordinate not in ws.merged_cells:
                    max_len = len(val_str)
            col_letter = col[0].column_letter
            ws.column_dimensions[col_letter].width = min(max(max_len + 4, 12), 60)

        buffer = io.BytesIO()
        wb.save(buffer)
        buffer.seek(0)
        file_content = buffer.getvalue()

        filename = f"audit_report_{tab}_{timezone.now().strftime('%Y%m%d_%H%M%S')}.xlsx"
        response = HttpResponse(
            file_content,
            content_type='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet'
        )
        response['Content-Disposition'] = f'attachment; filename="{filename}"'
        response['Content-Length'] = len(file_content)
        response['Cache-Control'] = 'private, no-store'
        return response
