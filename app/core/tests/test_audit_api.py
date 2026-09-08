import io

from auditlog.models import LogEntry
from django.contrib.auth import get_user_model
from django.contrib.contenttypes.models import ContentType
from django.test import RequestFactory, TestCase, override_settings
import openpyxl
from rest_framework import status
from rest_framework.test import APIClient

from core.models import UserActivityLog, UserManagementAudit, UserProfile, UserRole
from core.signals import get_client_ip
from core.audit_plant_resolver import resolve_plant_id
from tochuc.models import NhaMay

User = get_user_model()


class AuditApiAndSnapshotTestCase(TestCase):
    def setUp(self):
        self.client = APIClient()
        self.factory = RequestFactory()

        # Tạo 2 nhà máy thử nghiệm
        self.plant_vs = NhaMay.objects.create(ma_nha_may="VS", ten_nha_may="Nhà máy Vĩnh Sơn")
        self.plant_sh = NhaMay.objects.create(ma_nha_may="SH", ten_nha_may="Nhà máy Sông Hinh")

        # 1. Superuser
        self.admin_user = User.objects.create_superuser(
            username="admin_test",
            email="admin@test.com",
            password="Vsh@2026StrongPass!",
            first_name="Admin",
            last_name="Super"
        )
        UserProfile.objects.update_or_create(
            user=self.admin_user,
            defaults={"nha_may": self.plant_vs, "is_all_factories": True}
        )

        # 2. User chỉ có quyền xem nhật ký truy cập (Nhà máy Vĩnh Sơn)
        self.activity_user = User.objects.create_user(
            username="activity_user",
            email="activity@test.com",
            password="Vsh@2026StrongPass!",
            first_name="Hoạt Động",
            last_name="Vĩnh Sơn"
        )
        UserProfile.objects.update_or_create(
            user=self.activity_user,
            defaults={
                "nha_may": self.plant_vs,
                "is_all_factories": False,
                "can_view_activity_logs": True,
                "can_view_data_audit_logs": False,
                "can_view_user_management_audit": False,
            }
        )

        # 3. User chỉ có quyền xem vết thay đổi dữ liệu (Nhà máy Sông Hinh)
        self.data_user = User.objects.create_user(
            username="data_user",
            email="data@test.com",
            password="Vsh@2026StrongPass!",
            first_name="Kiểm Toán",
            last_name="Sông Hinh"
        )
        UserProfile.objects.update_or_create(
            user=self.data_user,
            defaults={
                "nha_may": self.plant_sh,
                "is_all_factories": False,
                "can_view_activity_logs": False,
                "can_view_data_audit_logs": True,
                "can_view_user_management_audit": False,
            }
        )

        # 4. User chỉ có is_all_factories nhưng KHÔNG có quyền xem log
        self.all_plants_no_perm_user = User.objects.create_user(
            username="all_no_perm",
            email="all_no_perm@test.com",
            password="Vsh@2026StrongPass!"
        )
        UserProfile.objects.update_or_create(
            user=self.all_plants_no_perm_user,
            defaults={
                "nha_may": self.plant_vs,
                "is_all_factories": True,
                "can_view_activity_logs": False,
                "can_view_data_audit_logs": False,
                "can_view_user_management_audit": False,
            }
        )
        self.all_plants_audit_user = User.objects.create_user(
            username="all_audit",
            email="all_audit@test.com",
            password="Vsh@2026StrongPass!",
        )
        UserProfile.objects.update_or_create(
            user=self.all_plants_audit_user,
            defaults={
                "nha_may": self.plant_vs,
                "is_all_factories": True,
                "can_view_activity_logs": True,
                "can_view_data_audit_logs": True,
                "can_view_user_management_audit": True,
                "can_export_audit_logs": True,
            },
        )
        # 5. User có quyền xuất Excel (Nhà máy Vĩnh Sơn)
        self.export_user = User.objects.create_user(
            username="export_user",
            email="export@test.com",
            password="Vsh@2026StrongPass!",
            first_name="Xuất",
            last_name="Báo Cáo"
        )
        UserProfile.objects.update_or_create(
            user=self.export_user,
            defaults={
                "nha_may": self.plant_vs,
                "is_all_factories": False,
                "can_view_activity_logs": True,
                "can_export_audit_logs": True,
            }
        )
        self.export_only_user = User.objects.create_user(
            username="export_only",
            email="export_only@test.com",
            password="Vsh@2026StrongPass!",
        )
        UserProfile.objects.update_or_create(
            user=self.export_only_user,
            defaults={
                "nha_may": self.plant_vs,
                "can_export_audit_logs": True,
            },
        )

        # Tạo mẫu UserActivityLog cho cả 2 nhà máy
        self.act_vs = UserActivityLog.objects.create(
            user=self.activity_user,
            action_type="LOGIN",
            description="Đăng nhập Vĩnh Sơn",
            ip_address="192.168.1.50",
            nha_may=self.plant_vs
        )
        self.act_sh = UserActivityLog.objects.create(
            user=self.data_user,
            action_type="LOGIN",
            description="Đăng nhập Sông Hinh",
            ip_address="192.168.2.60",
            nha_may=self.plant_sh
        )
        self.act_unscoped = UserActivityLog.objects.create(
            user=None,
            action_type="LOGIN_FAILED",
            description="Đăng nhập thất bại vô danh",
            ip_address="10.0.0.1",
            nha_may=None
        )

        # Tạo mẫu LogEntry với snapshot additional_data
        ct_nhamay = ContentType.objects.get_for_model(NhaMay)
        self.entry_vs = LogEntry.objects.create(
            content_type=ct_nhamay,
            object_pk=str(self.plant_vs.pk),
            object_repr="Nhà máy Vĩnh Sơn",
            action=LogEntry.Action.UPDATE,
            changes={"ten_nha_may": ["Cũ", "Nhà máy Vĩnh Sơn"]},
            actor=self.admin_user,
            additional_data={"plant_id": self.plant_vs.pk, "plant_code": "VS"}
        )
        self.entry_sh = LogEntry.objects.create(
            content_type=ct_nhamay,
            object_pk=str(self.plant_sh.pk),
            object_repr="Nhà máy Sông Hinh",
            action=LogEntry.Action.UPDATE,
            changes={"ten_nha_may": ["Cũ", "Nhà máy Sông Hinh"], "password": ["secret1", "secret2"]},
            actor=self.admin_user,
            additional_data={"plant_id": self.plant_sh.pk, "plant_code": "SH"}
        )
        self.entry_unscoped = LogEntry.objects.create(
            content_type=ct_nhamay,
            object_pk="999",
            object_repr="Bản ghi toàn cục",
            action=LogEntry.Action.CREATE,
            changes={"ten": ["A", "B"]},
            actor=self.admin_user,
            additional_data={}  # Không có plant_id
        )

        # Tạo mẫu UserManagementAudit với 4 mã chuẩn
        self.audit_create = UserManagementAudit.objects.create(
            actor=self.admin_user,
            target=self.activity_user,
            nha_may=self.plant_vs,
            action="create",
            changes={"role_id": 1}
        )
        self.audit_edit = UserManagementAudit.objects.create(
            actor=self.admin_user,
            target=self.activity_user,
            nha_may=self.plant_vs,
            action="edit",
            changes={"first_name": "Mới"}
        )
        self.audit_role = UserManagementAudit.objects.create(
            actor=self.admin_user,
            target=self.data_user,
            nha_may=self.plant_sh,
            action="role",
            changes={"role_id": 2}
        )
        self.audit_status = UserManagementAudit.objects.create(
            actor=self.admin_user,
            target=self.data_user,
            nha_may=self.plant_sh,
            action="status",
            changes={"is_active": False}
        )
        self.audit_unscoped = UserManagementAudit.objects.create(
            actor=self.admin_user,
            target=self.activity_user,
            nha_may=None,
            action="edit",
            changes={"email": ["old@test.com", "new@test.com"]},
        )

    # ================= 1. KIỂM TRA PHÂN QUYỀN RBAC =================
    def test_activity_log_permissions(self):
        """User có can_view_activity_logs truy cập 200, user không có nhận 403"""
        self.client.force_authenticate(user=self.activity_user)
        res = self.client.get('/api/audit/activity-logs/')
        self.assertEqual(res.status_code, status.HTTP_200_OK)

        self.client.force_authenticate(user=self.data_user)
        res = self.client.get('/api/audit/activity-logs/')
        self.assertEqual(res.status_code, status.HTTP_403_FORBIDDEN)

    def test_data_changes_permissions(self):
        """User có can_view_data_audit_logs truy cập 200, user không có nhận 403"""
        self.client.force_authenticate(user=self.data_user)
        res = self.client.get('/api/audit/data-changes/')
        self.assertEqual(res.status_code, status.HTTP_200_OK)

        self.client.force_authenticate(user=self.activity_user)
        res = self.client.get('/api/audit/data-changes/')
        self.assertEqual(res.status_code, status.HTTP_403_FORBIDDEN)

    def test_user_management_audit_permissions(self):
        """User không có quyền nhận 403, superuser nhận 200"""
        self.client.force_authenticate(user=self.activity_user)
        res = self.client.get('/api/audit/user-management/')
        self.assertEqual(res.status_code, status.HTTP_403_FORBIDDEN)

        self.client.force_authenticate(user=self.admin_user)
        res = self.client.get('/api/audit/user-management/')
        self.assertEqual(res.status_code, status.HTTP_200_OK)

    def test_is_all_factories_alone_does_not_grant_permission(self):
        """is_all_factories = True nhưng thiếu quyền xem tab vẫn bị chặn 403"""
        self.client.force_authenticate(user=self.all_plants_no_perm_user)
        res_act = self.client.get('/api/audit/activity-logs/')
        self.assertEqual(res_act.status_code, status.HTTP_403_FORBIDDEN)

        res_data = self.client.get('/api/audit/data-changes/')
        self.assertEqual(res_data.status_code, status.HTTP_403_FORBIDDEN)

        res_mgt = self.client.get('/api/audit/user-management/')
        self.assertEqual(res_mgt.status_code, status.HTTP_403_FORBIDDEN)

    # ================= 2. KIỂM TRA PHẠM VI NHÀ MÁY =================
    def test_plant_isolation_single_plant(self):
        """User thuộc Vĩnh Sơn chỉ thấy log Vĩnh Sơn, không thấy Sông Hinh hay unscoped"""
        self.client.force_authenticate(user=self.activity_user)
        res = self.client.get('/api/audit/activity-logs/')
        self.assertEqual(res.status_code, status.HTTP_200_OK)
        results = res.data.get('results', res.data)
        ids = [item['id'] for item in results]
        self.assertIn(self.act_vs.id, ids)
        self.assertNotIn(self.act_sh.id, ids)
        self.assertNotIn(self.act_unscoped.id, ids)

    def test_data_changes_plant_isolation(self):
        """User Sông Hinh chỉ thấy LogEntry có additional_data.plant_id của Sông Hinh"""
        self.client.force_authenticate(user=self.data_user)
        res = self.client.get('/api/audit/data-changes/')
        self.assertEqual(res.status_code, status.HTTP_200_OK)
        results = res.data.get('results', res.data)
        ids = [item['id'] for item in results]
        self.assertIn(self.entry_sh.id, ids)
        self.assertNotIn(self.entry_vs.id, ids)
        self.assertNotIn(self.entry_unscoped.id, ids)

    def test_superuser_sees_all_and_unscoped(self):
        """Superuser thấy toàn bộ log kể cả unscoped"""
        self.client.force_authenticate(user=self.admin_user)
        res = self.client.get('/api/audit/data-changes/')
        self.assertEqual(res.status_code, status.HTTP_200_OK)
        results = res.data.get('results', res.data)
        ids = [item['id'] for item in results]
        self.assertIn(self.entry_vs.id, ids)
        self.assertIn(self.entry_sh.id, ids)
        self.assertIn(self.entry_unscoped.id, ids)

    def test_all_factories_non_superuser_cannot_see_unscoped_logs(self):
        self.client.force_authenticate(user=self.all_plants_audit_user)
        activity = self.client.get('/api/audit/activity-logs/').data['results']
        management = self.client.get('/api/audit/user-management/').data['results']
        data_changes = self.client.get('/api/audit/data-changes/').data['results']
        self.assertNotIn(self.act_unscoped.id, [item['id'] for item in activity])
        self.assertNotIn(self.audit_unscoped.id, [item['id'] for item in management])
        self.assertNotIn(self.entry_unscoped.id, [item['id'] for item in data_changes])

    # ================= 3. KIỂM TRA MÃ HÀNH ĐỘNG & BẢO VỆ DỮ LIỆU NHẠY CẢM =================
    def test_user_management_action_display(self):
        """Kiểm tra ánh xạ 4 mã hành động create, edit, role, status sang tiếng Việt"""
        self.client.force_authenticate(user=self.admin_user)
        res = self.client.get('/api/audit/user-management/')
        self.assertEqual(res.status_code, status.HTTP_200_OK)
        results = res.data.get('results', res.data)
        action_map = {item['action']: item['action_display'] for item in results}
        self.assertEqual(action_map.get('create'), 'Tạo tài khoản')
        self.assertEqual(action_map.get('edit'), 'Sửa thông tin')
        self.assertEqual(action_map.get('role'), 'Đổi vai trò')
        self.assertEqual(action_map.get('status'), 'Khóa/Mở tài khoản')

    def test_sensitive_data_scrubbing(self):
        """Trường password trong changes phải được che giấu thành [ĐÃ ẨN]"""
        self.client.force_authenticate(user=self.admin_user)
        res = self.client.get(f'/api/audit/data-changes/{self.entry_sh.id}/')
        self.assertEqual(res.status_code, status.HTTP_200_OK)
        changes = res.data['changes']
        self.assertIn('password', changes)
        self.assertEqual(changes['password'], ['[ĐÃ ẨN]', '[ĐÃ ẨN]'])

    def test_nested_sensitive_data_and_free_text_are_scrubbed(self):
        self.entry_sh.changes = {
            "config": {
                "api_key": "top-secret",
                "note": "Bearer abc.def.ghi",
            }
        }
        self.entry_sh.object_repr = "token=plain-secret"
        self.entry_sh.save(update_fields=['changes', 'object_repr'])
        self.client.force_authenticate(user=self.admin_user)
        res = self.client.get(f'/api/audit/data-changes/{self.entry_sh.id}/')
        self.assertEqual(res.data['changes']['config']['api_key'], '[ĐÃ ẨN]')
        self.assertNotIn('abc.def.ghi', str(res.data))
        self.assertNotIn('plain-secret', res.data['object_repr'])

    def test_user_management_status_uses_new_value_from_diff(self):
        self.audit_status.changes = {'is_active': [True, False]}
        self.audit_status.save(update_fields=['changes'])
        self.client.force_authenticate(user=self.admin_user)
        res = self.client.get(f'/api/audit/user-management/{self.audit_status.id}/')
        self.assertIn('Bị khóa', res.data['changes_display'])

    # ================= 4. KIỂM TRA AN TOÀN IP & SPOOFING =================
    @override_settings(
        TRUST_PROXY_HEADERS=True,
        TRUSTED_PROXY_CIDRS=('127.0.0.1/32', '172.16.0.0/12'),
    )
    def test_get_client_ip_resolution(self):
        """get_client_ip trích xuất đúng IP từ header hợp lệ và fallback REMOTE_ADDR an toàn"""
        # 1. Fallback REMOTE_ADDR khi không có header
        req = self.factory.get('/', REMOTE_ADDR='10.10.10.10')
        self.assertEqual(get_client_ip(req), '10.10.10.10')

        # 2. Đọc Cloudflare connecting IP
        req = self.factory.get('/', REMOTE_ADDR='127.0.0.1', HTTP_CF_CONNECTING_IP='203.0.113.195')
        self.assertEqual(get_client_ip(req), '203.0.113.195')

        # 3. Đọc Nginx X-Real-IP
        req = self.factory.get('/', REMOTE_ADDR='127.0.0.1', HTTP_X_REAL_IP='198.51.100.22')
        self.assertEqual(get_client_ip(req), '198.51.100.22')

        # 4. Bỏ qua header có IP không hợp lệ
        req = self.factory.get('/', REMOTE_ADDR='10.10.10.10', HTTP_CF_CONNECTING_IP='invalid_ip_string')
        self.assertEqual(get_client_ip(req), '10.10.10.10')

        # 5. Direct, untrusted peer cannot spoof forwarding headers
        req = self.factory.get(
            '/',
            REMOTE_ADDR='198.51.100.99',
            HTTP_X_FORWARDED_FOR='1.2.3.4',
        )
        self.assertEqual(get_client_ip(req), '198.51.100.99')

    # ================= 5. KIỂM TRA METADATA ENDPOINT =================
    def test_metadata_endpoint(self):
        """Metadata trả về danh mục nhà máy và model trong phạm vi quyền của user"""
        self.client.force_authenticate(user=self.data_user)
        res = self.client.get('/api/audit/metadata/')
        self.assertEqual(res.status_code, status.HTTP_200_OK)
        data = res.data
        self.assertIn('plants', data)
        self.assertIn('models', data)
        self.assertIn('activity_actions', data)
        self.assertIn('data_change_actions', data)
        self.assertIn('user_management_actions', data)
        # User Sông Hinh chỉ thấy nhà máy Sông Hinh trong metadata
        plant_codes = [p['ma_nha_may'] for p in data['plants']]
        self.assertEqual(plant_codes, ['SH'])

    def test_metadata_requires_an_audit_view_permission(self):
        self.client.force_authenticate(user=self.all_plants_no_perm_user)
        res = self.client.get('/api/audit/metadata/')
        self.assertEqual(res.status_code, status.HTTP_403_FORBIDDEN)

    def test_invalid_filters_return_400(self):
        self.client.force_authenticate(user=self.data_user)
        self.assertEqual(
            self.client.get('/api/audit/data-changes/?action=invalid').status_code,
            status.HTTP_400_BAD_REQUEST,
        )
        self.assertEqual(
            self.client.get('/api/audit/data-changes/?from_date=not-a-date').status_code,
            status.HTTP_400_BAD_REQUEST,
        )
        self.assertEqual(
            self.client.get(
                '/api/audit/data-changes/?from_date=2024-01-01&to_date=2026-01-02'
            ).status_code,
            status.HTTP_400_BAD_REQUEST,
        )

    def test_plant_name_is_loaded_from_database_not_fixed_id(self):
        self.plant_sh.ten_nha_may = 'Tên nhà máy thử nghiệm'
        self.plant_sh.save(update_fields=['ten_nha_may'])
        self.client.force_authenticate(user=self.data_user)
        res = self.client.get(f'/api/audit/data-changes/{self.entry_sh.id}/')
        self.assertEqual(res.data['plant_name'], 'Tên nhà máy thử nghiệm')

    def test_plant_resolver_fails_closed_for_non_model_values(self):
        self.assertIsNone(resolve_plant_id('not-a-model-instance'))

    # ================= 6. KIỂM TRA EXCEL EXPORT API & PHÒNG CHỐNG FORMULA INJECTION =================
    def test_export_excel_permission_denied(self):
        """Chặn 401 khi chưa đăng nhập và 403 khi không có quyền can_export_audit_logs"""
        # 1. Chưa đăng nhập -> 401
        res = self.client.post('/api/audit/export-excel/', {'tab': 'activity'})
        self.assertEqual(res.status_code, status.HTTP_401_UNAUTHORIZED)

        # 2. Đã đăng nhập nhưng thiếu quyền can_export_audit_logs -> 403
        self.client.force_authenticate(user=self.activity_user)
        res = self.client.post('/api/audit/export-excel/', {'tab': 'activity'})
        self.assertEqual(res.status_code, status.HTTP_403_FORBIDDEN)

        # Quyền xuất không tự cấp quyền xem nội dung của từng tab.
        self.client.force_authenticate(user=self.export_only_user)
        res = self.client.post('/api/audit/export-excel/', {'tab': 'activity'})
        self.assertEqual(res.status_code, status.HTTP_403_FORBIDDEN)

    def test_export_excel_activity_tab_success(self):
        """User có quyền xuất Excel tải thành công file .xlsx nhật ký truy cập kèm phân quyền nhà máy"""
        self.client.force_authenticate(user=self.export_user)
        res = self.client.post('/api/audit/export-excel/', {'tab': 'activity'})
        self.assertEqual(res.status_code, status.HTTP_200_OK)
        self.assertEqual(
            res['Content-Type'],
            'application/vnd.openxmlformats-officedocument.spreadsheetml.sheet'
        )
        self.assertIn('attachment; filename="audit_report_activity_', res['Content-Disposition'])

        wb = openpyxl.load_workbook(io.BytesIO(res.content))
        ws = wb.active
        self.assertEqual(ws.title, "Nhật ký truy cập")
        headers = [cell.value for cell in ws[1]]
        self.assertIn("STT", headers)
        self.assertIn("Tên đăng nhập", headers)
        self.assertIn("Địa chỉ IP", headers)

        # Do export_user thuộc nhà máy Vĩnh Sơn, chỉ thấy bản ghi Vĩnh Sơn
        usernames = [ws.cell(row=r, column=3).value for r in range(2, ws.max_row + 1)]
        self.assertIn(self.activity_user.username, usernames)
        self.assertNotIn(self.data_user.username, usernames)

    def test_export_excel_data_changes_tab_success(self):
        """Superuser xuất thành công tab vết thay đổi dữ liệu"""
        self.client.force_authenticate(user=self.admin_user)
        res = self.client.post('/api/audit/export-excel/', {'tab': 'data_changes'})
        self.assertEqual(res.status_code, status.HTTP_200_OK)
        wb = openpyxl.load_workbook(io.BytesIO(res.content))
        ws = wb.active
        self.assertEqual(ws.title, "Vết thay đổi dữ liệu")
        self.assertGreaterEqual(ws.max_row, 2)

    def test_export_excel_ignores_blank_optional_integer_filters(self):
        """Select chưa chọn không làm IntegerField trả lỗi HTTP 400."""
        self.client.force_authenticate(user=self.admin_user)
        res = self.client.post(
            '/api/audit/export-excel/',
            {
                'tab': 'data_changes',
                'plant_id': '',
                'content_type_id': '',
                'action': '',
            },
        )
        self.assertEqual(res.status_code, status.HTTP_200_OK)

    def test_export_excel_user_management_tab_success(self):
        """Superuser xuất thành công tab nhật ký quản trị tài khoản"""
        self.client.force_authenticate(user=self.admin_user)
        res = self.client.post('/api/audit/export-excel/', {'tab': 'user_management'})
        self.assertEqual(res.status_code, status.HTTP_200_OK)
        wb = openpyxl.load_workbook(io.BytesIO(res.content))
        ws = wb.active
        self.assertEqual(ws.title, "Quản trị tài khoản")
        self.assertGreaterEqual(ws.max_row, 2)

    def test_export_excel_invalid_tab_returns_400(self):
        """Gửi tab không hợp lệ trả về HTTP 400 Bad Request"""
        self.client.force_authenticate(user=self.admin_user)
        res = self.client.post('/api/audit/export-excel/', {'tab': 'invalid_unknown_tab'})
        self.assertEqual(res.status_code, status.HTTP_400_BAD_REQUEST)
        self.assertIn('tab', res.data)

    def test_export_excel_formula_injection_sanitization(self):
        """Phòng chống Formula Injection: các chuỗi bắt đầu bằng =, +, -, @, \\t, \\r được thêm tiền tố nháy đơn"""
        malicious_payloads = [
            "=cmd|' /C calc'!A0",
            "+SUM(A1:A10)",
            "-10*20",
            "@HYPERLINK('http://evil.com')",
            "\tTAB_INJECTION",
            "   =TRIMMED_FORMULA"
        ]
        for payload in malicious_payloads:
            UserActivityLog.objects.create(
                user=self.admin_user,
                action_type="LOGIN",
                description=payload,
                nha_may=self.plant_vs
            )

        self.client.force_authenticate(user=self.admin_user)
        res = self.client.post('/api/audit/export-excel/', {'tab': 'activity'})
        self.assertEqual(res.status_code, status.HTTP_200_OK)

        wb = openpyxl.load_workbook(io.BytesIO(res.content))
        ws = wb.active

        # Cột mô tả là cột 9
        descriptions = [str(ws.cell(row=r, column=9).value or '') for r in range(2, ws.max_row + 1)]

        for payload in malicious_payloads:
            sanitized_expected = "'" + payload
            self.assertIn(sanitized_expected, descriptions)
