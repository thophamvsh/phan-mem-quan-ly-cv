import io

import openpyxl
from django.contrib.auth import get_user_model
from django.core.files.uploadedfile import SimpleUploadedFile
from django.http import JsonResponse
from django.test import RequestFactory
from django.test import TestCase
from django.utils import timezone
from rest_framework.test import APIClient

from core.models import DataSyncAudit, UserProfile
from core.sync_audit import audit_excel_import, record_data_sync
from core.tasks import serialize_model_instance
from tochuc.models import NhaMay


class DataSyncAuditTests(TestCase):
    def setUp(self):
        self.client = APIClient()
        self.sh = NhaMay.objects.create(ma_nha_may="SH", ten_nha_may="Sông Hinh")
        self.vs = NhaMay.objects.create(ma_nha_may="VS", ten_nha_may="Vĩnh Sơn")
        user_model = get_user_model()
        self.user = user_model.objects.create_user(
            email="audit-sh@example.com", username="audit-sh", password="StrongPass123!"
        )
        UserProfile.objects.create(
            user=self.user, nha_may=self.sh, can_view_data_audit_logs=True,
            can_export_audit_logs=True,
        )
        self.admin = user_model.objects.create_superuser(
            email="audit-admin@example.com", username="audit-admin", password="StrongPass123!"
        )
        self.sh_log = record_data_sync(
            actor=self.user, nha_may=self.sh, source=DataSyncAudit.Source.EXCEL,
            data_type="Danh mục nhân sự", status=DataSyncAudit.Status.SUCCESS,
            started_at=timezone.now(), processed_count=3, created_count=2,
            updated_count=1, filename="../nhan-su.xlsx",
            error_summary="token=raw-secret",
            metadata={"token": "secret", "sheet": "Nhân sự"},
        )
        self.vs_log = record_data_sync(
            actor=self.admin, nha_may=self.vs, source=DataSyncAudit.Source.GOOGLE_SHEET,
            data_type="Sản lượng", status=DataSyncAudit.Status.FAILED,
            started_at=timezone.now(), processed_count=2, failed_count=2,
            error_summary="Không thể đồng bộ",
        )

    def test_record_is_sanitized_and_filename_has_no_path(self):
        self.assertEqual(self.sh_log.filename, "nhan-su.xlsx")
        self.assertEqual(self.sh_log.metadata["token"], "[ĐÃ ẨN]")
        self.assertNotIn("raw-secret", self.sh_log.error_summary)
        serialized = serialize_model_instance(self.sh_log)
        self.assertEqual(serialized["id"], str(self.sh_log.id))

    def test_single_plant_user_only_sees_own_plant(self):
        self.client.force_authenticate(self.user)
        response = self.client.get("/api/audit/data-syncs/")
        self.assertEqual(response.status_code, 200)
        ids = {item["id"] for item in response.data["results"]}
        self.assertIn(str(self.sh_log.id), ids)
        self.assertNotIn(str(self.vs_log.id), ids)

    def test_superuser_can_filter_status_and_plant(self):
        self.client.force_authenticate(self.admin)
        response = self.client.get(
            "/api/audit/data-syncs/", {"plant_id": self.vs.id, "action": "FAILED"}
        )
        self.assertEqual(response.status_code, 200)
        self.assertEqual([item["id"] for item in response.data["results"]], [str(self.vs_log.id)])

    def test_sync_export_is_scoped_and_valid_xlsx(self):
        self.client.force_authenticate(self.user)
        response = self.client.post("/api/audit/export-excel/", {"tab": "sync"})
        self.assertEqual(response.status_code, 200)
        workbook = openpyxl.load_workbook(io.BytesIO(response.content))
        sheet = workbook.active
        self.assertEqual(sheet.title, "Phiên đồng bộ")
        values = [sheet.cell(row=row, column=6).value for row in range(2, sheet.max_row + 1)]
        self.assertIn("Danh mục nhân sự", values)
        self.assertNotIn("Sản lượng", values)

    def test_excel_import_decorator_records_success_and_resets_file(self):
        upload = SimpleUploadedFile(
            "folder/devices.xlsx", b"fake-excel-content",
            content_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        )
        request = RequestFactory().post(
            "/fake-import/", {"file": upload, "factory_code": "SH"}
        )
        request.user = self.user

        @audit_excel_import("Danh mục thiết bị")
        def fake_import(incoming_request):
            self.assertEqual(incoming_request.FILES["file"].read(), b"fake-excel-content")
            return JsonResponse({"imported_count": 4})

        response = fake_import(request)

        self.assertEqual(response.status_code, 200)
        audit = DataSyncAudit.objects.filter(data_type="Danh mục thiết bị").get()
        self.assertEqual(audit.nha_may, self.sh)
        self.assertEqual(audit.processed_count, 4)
        self.assertEqual(audit.status, DataSyncAudit.Status.SUCCESS)
        self.assertEqual(len(audit.checksum_sha256), 64)
