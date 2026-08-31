from datetime import date, datetime, time
from io import BytesIO

from django.core.exceptions import ValidationError
from django.test import TestCase
from rest_framework.test import APITestCase

from core.models import User, UserProfile, UserRole
from tochuc.models import NhaMay
from nhatkyvanhanh.models import SogiaonhancaHC, SogiaonhancaVH
from quanlycatruc.models import BoPhan, ChiTietPhuongAnPhanCongCa, DieuChinhNhanSuCaTruc, DonViToChuc, KipTruc, LichTrucCa, MauChuKyCaTruc, NhanSu, NhomLichTruc, PhuongAnPhanCongCa, ThanhVienKipTruc
from quanlycatruc.services import abbreviated_staff_names, actual_shift_staff, generate_monthly_schedule, transition_schedule


def create_user(username, plant, **permissions):
    user = User.objects.create_user(username=username, email=f"{username}@example.com", password="test-pass-123")
    profile, _ = UserProfile.objects.get_or_create(user=user)
    profile.nha_may = plant
    for field, value in permissions.items():
        setattr(profile, field, value)
    profile.save()
    return user


class StaffNameDisplayTests(TestCase):
    def test_abbreviates_names_and_expands_collisions(self):
        labels = abbreviated_staff_names([
            "Phạm Hoàng Thọ", "Nguyễn Tá Quỳnh", "Phan Minh Thọ",
        ])
        self.assertEqual(labels["Phạm Hoàng Thọ"], "P.H.Thọ")
        self.assertEqual(labels["Nguyễn Tá Quỳnh"], "N.Quỳnh")
        self.assertEqual(labels["Phan Minh Thọ"], "P.M.Thọ")


class ScheduleFixtureMixin:
    def setUp(self):
        self.plant = NhaMay.objects.create(ma_nha_may="SH", ten_nha_may="Sông Hinh")
        self.other_plant = NhaMay.objects.create(ma_nha_may="VS", ten_nha_may="Vĩnh Sơn")
        for index, code in enumerate(("A", "B", "C", "D"), start=1):
            KipTruc.objects.create(nha_may=self.plant, loai_kip="van_hanh", ma_kip=code, ten_kip=f"Kíp {code}", thu_tu=index)
        for index, code in enumerate(("HC1", "HC2"), start=1):
            KipTruc.objects.create(nha_may=self.plant, loai_kip="hanh_chinh", ma_kip=code, ten_kip=f"Ca {index}", thu_tu=index)
        self.template = MauChuKyCaTruc.objects.create(
            nha_may=self.plant,
            ten_mau="Chu kỳ chuẩn",
            ngay_moc_van_hanh=date(2026, 8, 1),
            vi_tri_moc_van_hanh=0,
            ngay_moc_hanh_chinh=date(2026, 8, 1),
            vi_tri_moc_hanh_chinh=0,
            tu_ngay=date(2026, 1, 1),
        )
        self.user = create_user(
            "scheduler", self.plant,
            can_view_shift_schedule=True,
            can_create_shift_schedule=True,
            can_edit_shift_schedule=True,
            can_delete_shift_schedule=True,
            can_submit_shift_schedule=True,
            can_approve_shift_schedule=True,
            can_manage_shift_roster=True,
            can_export_shift_schedule=True,
        )

    def make_schedule(self, month=8, year=2026):
        return LichTrucCa.objects.create(nha_may=self.plant, thang=month, nam=year, mau_chu_ky=self.template, nguoi_tao=self.user)


class ScheduleServiceTests(ScheduleFixtureMixin, TestCase):
    def test_generates_operation_and_admin_cycle(self):
        schedule = self.make_schedule()
        generate_monthly_schedule(schedule.id, self.user)
        days = list(schedule.danh_sach_ngay.select_related("kip_ca_ngay", "kip_ca_dem", "kip_hanh_chinh", "kip_hanh_chinh_truoc"))
        self.assertEqual(len(days), 31)
        self.assertEqual(days[0].hien_thi_ca_ngay, "A/D")
        self.assertEqual(days[0].kip_ca_dem.ma_kip, "A")
        self.assertEqual(days[0].kip_ca_ngay.ma_kip, "D")
        self.assertEqual(days[3].hien_thi_ca_ngay, "D/B")
        self.assertEqual(days[3].hien_thi_ca_hanh_chinh, "2/1")
        self.assertEqual(days[9].hien_thi_ca_hanh_chinh, "1/2")
        self.assertEqual([day.kip_hanh_chinh.ma_kip for day in days[:12]].count("HC1"), 6)
        self.assertEqual([day.kip_hanh_chinh.ma_kip for day in days[:12]].count("HC2"), 6)

    def test_generates_from_schedule_specific_cycle_anchor(self):
        schedule = LichTrucCa.objects.create(
            nha_may=self.plant,
            thang=9,
            nam=2026,
            mau_chu_ky=self.template,
            nguoi_tao=self.user,
            che_do_sinh_chu_ky=LichTrucCa.CheDoSinhChuKy.MOC_TUY_CHON,
            ngay_moc_chu_ky=date(2026, 9, 3),
            vi_tri_moc_van_hanh=0,
            vi_tri_moc_hanh_chinh=3,
        )
        generate_monthly_schedule(schedule.id, self.user)
        days = list(schedule.danh_sach_ngay.select_related(
            "kip_ca_ngay", "kip_ca_dem", "kip_hanh_chinh"
        ).order_by("ngay")[:4])
        # Ngày 03/09 đúng vị trí mốc đã chọn; các ngày trước/sau được suy ra theo chu kỳ 12 ngày.
        self.assertEqual((days[2].kip_ca_ngay.ma_kip, days[2].kip_ca_dem.ma_kip), ("D", "A"))
        self.assertTrue(days[2].la_ngay_chuyen_kip)
        self.assertEqual(days[2].kip_hanh_chinh.ma_kip, "HC1")
        self.assertTrue(days[2].la_ngay_chuyen_ca_hc)
        self.assertEqual((days[3].kip_ca_ngay.ma_kip, days[3].kip_ca_dem.ma_kip), ("D", "A"))

    def test_cycle_continues_across_month_boundary(self):
        august = self.make_schedule()
        september = self.make_schedule(month=9)
        generate_monthly_schedule(august.id, self.user)
        generate_monthly_schedule(september.id, self.user)
        aug_last = august.danh_sach_ngay.select_related("kip_ca_ngay", "kip_ca_dem").last()
        sep_first = september.danh_sach_ngay.select_related("kip_ca_ngay", "kip_ca_dem").first()
        # 31/08 is position 7 (zero-based 6), 01/09 is the following position 8.
        self.assertEqual((aug_last.kip_ca_ngay.ma_kip, aug_last.kip_ca_dem.ma_kip), ("C", "B"))
        self.assertEqual((sep_first.kip_ca_ngay.ma_kip, sep_first.kip_ca_dem.ma_kip), ("C", "B"))
        self.assertFalse(sep_first.la_ngay_chuyen_kip)

    def test_new_month_continues_from_corrected_locked_previous_month(self):
        august = self.make_schedule()
        generate_monthly_schedule(august.id, self.user)
        teams = {team.ma_kip: team for team in KipTruc.objects.filter(nha_may=self.plant)}
        last_days = list(august.danh_sach_ngay.order_by("-ngay")[:3])[::-1]
        for index, day in enumerate(last_days):
            day.kip_ca_ngay = teams["A"]
            day.kip_ca_dem = teams["C"]
            day.la_ngay_chuyen_kip = index == 0
            day.save()
        transition_schedule(august.id, self.user, "gui_duyet")
        transition_schedule(august.id, self.user, "phe_duyet")
        transition_schedule(august.id, self.user, "khoa")

        september = self.make_schedule(month=9)
        generate_monthly_schedule(september.id, self.user)
        sep_first = september.danh_sach_ngay.select_related("kip_ca_ngay", "kip_ca_dem").first()

        self.assertEqual((sep_first.kip_ca_ngay.ma_kip, sep_first.kip_ca_dem.ma_kip), ("C", "B"))

    def test_new_month_continues_from_previous_month_currently_in_use(self):
        august = self.make_schedule()
        generate_monthly_schedule(august.id, self.user)
        teams = {team.ma_kip: team for team in KipTruc.objects.filter(nha_may=self.plant)}
        last_days = list(august.danh_sach_ngay.order_by("-ngay")[:12])[::-1]
        operation_cycle = [
            ("B", "A"), ("B", "A"), ("B", "A"),
            ("C", "B"), ("C", "B"), ("C", "B"),
            ("D", "C"), ("D", "C"), ("D", "C"),
            ("A", "D"), ("A", "D"), ("A", "D"),
        ]
        for index, (day, (day_team, night_team)) in enumerate(zip(last_days, operation_cycle)):
            day.kip_ca_ngay = teams[day_team]
            day.kip_ca_dem = teams[night_team]
            day.kip_hanh_chinh = teams["HC2"]
            day.la_ngay_chuyen_kip = index % 3 == 0
            day.save()
        transition_schedule(august.id, self.user, "gui_duyet")
        transition_schedule(august.id, self.user, "phe_duyet")
        transition_schedule(august.id, self.user, "ap_dung")

        september = self.make_schedule(month=9)
        generate_monthly_schedule(september.id, self.user)
        sep_first = september.danh_sach_ngay.select_related(
            "kip_ca_ngay", "kip_ca_dem", "kip_hanh_chinh"
        ).first()

        self.assertEqual(
            (sep_first.kip_ca_ngay.ma_kip, sep_first.kip_ca_dem.ma_kip),
            ("B", "A"),
        )
        self.assertTrue(sep_first.la_ngay_chuyen_kip)
        self.assertEqual(sep_first.kip_hanh_chinh.ma_kip, "HC2")

    def test_cannot_regenerate_submitted_schedule(self):
        schedule = self.make_schedule()
        generate_monthly_schedule(schedule.id, self.user)
        transition_schedule(schedule.id, self.user, "gui_duyet")
        with self.assertRaises(ValidationError):
            generate_monthly_schedule(schedule.id, self.user)

    def test_approval_creates_snapshot_and_locks_editing(self):
        schedule = self.make_schedule()
        generate_monthly_schedule(schedule.id, self.user)
        transition_schedule(schedule.id, self.user, "gui_duyet")
        approved = transition_schedule(schedule.id, self.user, "phe_duyet")
        self.assertEqual(approved.trang_thai, LichTrucCa.TrangThai.DA_DUYET)
        self.assertIsNotNone(approved.ngay_duyet)
        self.assertFalse(approved.co_the_chinh_sua)

    def test_team_code_can_follow_each_plants_operating_model(self):
        team = KipTruc.objects.create(
            nha_may=self.other_plant,
            loai_kip="van_hanh",
            ma_kip="DAP1",
            ten_kip="Kíp trực đập 1",
        )
        self.assertEqual(team.ma_kip, "DAP1")


class ScheduleApiTests(ScheduleFixtureMixin, APITestCase):
    def setUp(self):
        super().setUp()
        self.client.force_authenticate(self.user)

    def _create_group(self, code="VHNM"):
        unit = DonViToChuc.objects.create(
            ma_don_vi=f"QL{code}", ten_don_vi="Quản lý vận hành",
            loai_don_vi="nha_may", nha_may=self.plant,
        )
        department = BoPhan.objects.create(
            don_vi=unit, ma_bo_phan=f"BP{code}", ten_bo_phan="Vận hành nhà máy",
            loai_bo_phan="van_hanh",
        )
        group = NhomLichTruc.objects.create(
            nha_may=self.plant, don_vi=unit, bo_phan=department,
            ma_nhom=code, ten_nhom="Lịch vận hành nhà máy",
        )
        return group, unit, department

    def test_individual_grant_allows_access_when_role_denies_permission(self):
        role = UserRole.objects.create(
            name="Nhân viên không có quyền lịch trực",
            permissions={"can_view_shift_schedule": False},
        )
        profile = self.user.profile
        profile.role = role
        profile.individual_permissions = {"can_view_shift_schedule": True}
        profile.save()

        response = self.client.get(
            f"/api/v1/quanlycatruc/lich-truc/?nha_may={self.plant.id}"
        )

        self.assertEqual(response.status_code, 200, response.data)

    def test_can_update_and_delete_unused_shift_group(self):
        group, _, _ = self._create_group("EDIT")

        updated = self.client.patch(
            f"/api/v1/quanlycatruc/nhom-lich/{group.id}/",
            {"ten_nhom": "Lịch vận hành đã sửa", "dia_diem": "Gian máy"},
            format="json",
        )
        self.assertEqual(updated.status_code, 200, updated.data)
        self.assertEqual(updated.data["ten_nhom"], "Lịch vận hành đã sửa")

        deleted = self.client.delete(f"/api/v1/quanlycatruc/nhom-lich/{group.id}/")
        self.assertEqual(deleted.status_code, 204, deleted.data)
        self.assertFalse(NhomLichTruc.objects.filter(id=group.id).exists())

    def test_cannot_delete_shift_group_with_related_data(self):
        group, _, _ = self._create_group("USED")
        KipTruc.objects.create(
            nha_may=self.plant, nhom_lich=group, loai_kip="van_hanh",
            ma_kip="USED-A", ten_kip="Kíp đang sử dụng",
        )

        deleted = self.client.delete(f"/api/v1/quanlycatruc/nhom-lich/{group.id}/")

        self.assertEqual(deleted.status_code, 400, deleted.data)
        self.assertIn("kíp trực", str(deleted.data))
        self.assertTrue(NhomLichTruc.objects.filter(id=group.id).exists())

    def test_group_scopes_teams_and_personnel_options(self):
        unit = DonViToChuc.objects.create(
            ma_don_vi="QLVH", ten_don_vi="Quản lý vận hành",
            loai_don_vi="nha_may", nha_may=self.plant,
        )
        department = BoPhan.objects.create(
            don_vi=unit, ma_bo_phan="VHNM", ten_bo_phan="Vận hành nhà máy",
            loai_bo_phan="van_hanh",
        )
        person = NhanSu.objects.create(
            ho_ten="Nguyễn Văn A", don_vi=unit, bo_phan=department,
            chuc_danh="Trực chính",
        )
        group_response = self.client.post("/api/v1/quanlycatruc/nhom-lich/", {
            "nha_may": self.plant.id, "don_vi": unit.id, "bo_phan": department.id,
            "ma_nhom": "VHNM", "ten_nhom": "Lịch vận hành nhà máy",
            "dia_diem": "Nhà máy Sông Hinh",
        }, format="json")
        self.assertEqual(group_response.status_code, 201, group_response.data)
        group_id = group_response.data["id"]
        scope_response = self.client.post("/api/v1/quanlycatruc/pham-vi-nhan-su/", {
            "nhom_lich": group_id, "nhan_su": person.id, "tu_ngay": "2026-08-01",
        }, format="json")
        self.assertEqual(scope_response.status_code, 201, scope_response.data)
        repeated_scope = self.client.post("/api/v1/quanlycatruc/pham-vi-nhan-su/", {
            "nhom_lich": group_id, "nhan_su": person.id, "tu_ngay": "2026-08-01",
        }, format="json")
        self.assertEqual(repeated_scope.status_code, 201, repeated_scope.data)
        self.assertEqual(
            NhomLichTruc.objects.get(id=group_id).pham_vi_nhan_su.filter(nhan_su=person).count(),
            1,
        )
        scopes = self.client.get(
            f"/api/v1/quanlycatruc/pham-vi-nhan-su/?nhom_lich={group_id}"
        )
        self.assertEqual(scopes.status_code, 200, scopes.data)
        self.assertIsInstance(scopes.data, list)
        self.assertEqual([item["nhan_su"] for item in scopes.data], [person.id])
        team_response = self.client.post("/api/v1/quanlycatruc/kip-truc/", {
            "nha_may": self.plant.id, "nhom_lich": group_id,
            "loai_kip": "van_hanh", "ma_kip": "DAP1", "ten_kip": "Kíp trực đập 1",
        }, format="json")
        self.assertEqual(team_response.status_code, 201, team_response.data)
        options = self.client.get(f"/api/v1/quanlycatruc/kip-truc/nhan-su-options/?nha_may={self.plant.id}&nhom_lich={group_id}")
        self.assertEqual(options.status_code, 200, options.data)
        self.assertEqual([item["id"] for item in options.data], [person.id])

    def test_shift_handover_can_load_actual_roster_and_store_schedule_source(self):
        profile = self.user.profile
        profile.can_view_shift_handover_logs = True
        profile.can_create_shift_handover_logs = True
        profile.can_edit_own_shift_handover_logs = True
        profile.save(update_fields=[
            "can_view_shift_handover_logs",
            "can_create_shift_handover_logs",
            "can_edit_own_shift_handover_logs",
        ])
        unit = DonViToChuc.objects.create(
            ma_don_vi="SH-GC", ten_don_vi="Vận hành giao ca",
            loai_don_vi="nha_may", nha_may=self.plant,
        )
        department = BoPhan.objects.create(
            don_vi=unit, ma_bo_phan="VH-GC", ten_bo_phan="Vận hành",
            loai_bo_phan="van_hanh",
        )
        team_a = KipTruc.objects.get(nha_may=self.plant, ma_kip="A")
        ThanhVienKipTruc.objects.create(
            kip_truc=team_a, user=self.user, vai_tro="truong_ca",
            tu_ngay=date(2026, 1, 1),
        )
        for index, (name, role) in enumerate((
            ("Nhân sự trực chính", "truc_chinh"),
            ("Nhân sự trực phụ", "truc_phu"),
            ("Kỹ thuật viên không đưa vào sổ", "ky_thuat_vien"),
        ), start=1):
            person = NhanSu.objects.create(
                ma_nhan_vien=f"GC{index}", ho_ten=name,
                don_vi=unit, bo_phan=department,
            )
            ThanhVienKipTruc.objects.create(
                kip_truc=team_a, nhan_su=person, vai_tro=role,
                tu_ngay=date(2026, 1, 1),
            )
        schedule = self.make_schedule()
        generate_monthly_schedule(schedule.id, self.user)
        transition_schedule(schedule.id, self.user, "gui_duyet")
        transition_schedule(schedule.id, self.user, "phe_duyet")

        roster = self.client.get(
            "/api/nhatkyvanhanh/so-giao-nhan-ca-vh/bien-che-lich-truc/",
            {
                "nha_may": self.plant.id,
                "ngay_truc": "2026-08-01",
                "ca_truc": "A",
                "loai_thoi_gian_truc": "dem",
            },
        )
        self.assertEqual(roster.status_code, 200, roster.data)
        self.assertTrue(roster.data["truong_ca_khop_nguoi_tao"])
        self.assertEqual(
            [(item["ho_ten"], item["vai_tro"]) for item in roster.data["nhan_su"]],
            [("Nhân sự trực chính", "truc_chinh"), ("Nhân sự trực phụ", "truc_phu")],
        )

        created = self.client.post(
            "/api/nhatkyvanhanh/so-giao-nhan-ca-vh/",
            {
                "ngay_truc": "2026-08-01",
                "ca_truc": "A",
                "loai_thoi_gian_truc": "dem",
                "thoi_gian_bat_dau_ca": "2026-08-01T20:00:00+07:00",
                "thoi_gian_giao_ca": "2026-08-02T08:00:00+07:00",
                "lich_truc_nguon": roster.data["lich_truc_id"],
                "ngay_truc_ca_nguon": roster.data["ngay_truc_id"],
                "phien_ban_lich_nguon": 999,
            },
            format="json",
        )
        self.assertEqual(created.status_code, 201, created.data)
        shift_log = SogiaonhancaVH.objects.get(pk=created.data["id"])
        self.assertEqual(shift_log.lich_truc_nguon_id, schedule.id)
        self.assertEqual(shift_log.phien_ban_lich_nguon, schedule.phien_ban)
        self.assertIsNotNone(shift_log.dong_bo_bien_che_at)

    def test_shift_handover_uses_the_only_scheduled_team_when_form_has_stale_default(self):
        schedule = self.make_schedule()
        generate_monthly_schedule(schedule.id, self.user)
        transition_schedule(schedule.id, self.user, "gui_duyet")
        transition_schedule(schedule.id, self.user, "phe_duyet")
        last_day = schedule.danh_sach_ngay.select_related("kip_ca_ngay").get(
            ngay=date(2026, 8, 31)
        )
        scheduled_shift = last_day.kip_ca_ngay.ma_kip
        stale_form_shift = next(code for code in ("A", "B", "C", "D") if code != scheduled_shift)

        roster = self.client.get(
            "/api/nhatkyvanhanh/so-giao-nhan-ca-vh/bien-che-lich-truc/",
            {
                "nha_may": self.plant.id,
                "ngay_truc": "2026-08-31",
                "ca_truc": stale_form_shift,
                "loai_thoi_gian_truc": "ngay",
            },
        )

        self.assertEqual(roster.status_code, 200, roster.data)
        self.assertEqual(roster.data["ca_truc"], scheduled_shift)
        self.assertTrue(roster.data["ca_truc_tu_dong_dieu_chinh"])
        self.assertEqual(roster.data["ngay_truc_id"], last_day.id)

    def test_admin_shift_handover_can_load_hc_roster_without_requiring_creator_membership(self):
        profile = self.user.profile
        profile.can_view_admin_shift_handover_logs = True
        profile.can_create_admin_shift_handover_logs = True
        profile.save(update_fields=[
            "can_view_admin_shift_handover_logs",
            "can_create_admin_shift_handover_logs",
        ])
        unit = DonViToChuc.objects.create(
            ma_don_vi="SH-HC", ten_don_vi="Hành chính", loai_don_vi="nha_may", nha_may=self.plant,
        )
        department = BoPhan.objects.create(
            don_vi=unit, ma_bo_phan="HC", ten_bo_phan="Kỹ thuật hành chính", loai_bo_phan="hanh_chinh",
        )
        person = NhanSu.objects.create(
            ma_nhan_vien="HC001", ho_ten="Nhân sự KT-HC", don_vi=unit, bo_phan=department,
        )
        first_admin_team_code = self.template.chu_ky_hanh_chinh[0]["team"]
        hc1 = KipTruc.objects.get(nha_may=self.plant, ma_kip=first_admin_team_code)
        ThanhVienKipTruc.objects.create(
            kip_truc=hc1, nhan_su=person, vai_tro="nhan_vien", tu_ngay=date(2026, 1, 1),
        )
        schedule = self.make_schedule()
        generate_monthly_schedule(schedule.id, self.user)
        transition_schedule(schedule.id, self.user, "gui_duyet")
        transition_schedule(schedule.id, self.user, "phe_duyet")

        roster = self.client.get(
            "/api/nhatkyvanhanh/so-giao-nhan-ca-hc/bien-che-lich-truc/",
            {
                "nha_may": self.plant.id,
                "ngay_truc": "2026-08-01",
                "thoi_gian_bat_dau": "2026-08-01T08:00:00+07:00",
                "thoi_gian_ket_thuc": "2026-08-01T17:00:00+07:00",
            },
        )
        self.assertEqual(roster.status_code, 200, roster.data)
        selection = roster.data["lua_chon"][0]
        self.assertEqual([item["ho_ten"] for item in selection["nhan_su"]], ["Nhân sự KT-HC"])
        self.assertIn("thoi_gian_bat_dau", selection["nhan_su"][0])
        self.assertIn("thoi_gian_ket_thuc", selection["nhan_su"][0])
        self.assertFalse(selection["nguoi_tao_thuoc_bien_che"])

        created = self.client.post(
            "/api/nhatkyvanhanh/so-giao-nhan-ca-hc/",
            {
                "nha_may": self.plant.id,
                "ngay_truc": "2026-08-01",
                "thoi_gian_bat_dau_ca": "2026-08-01T08:00:00+07:00",
                "thoi_gian_giao_ca": "2026-08-01T17:00:00+07:00",
                "lich_truc_nguon": selection["lich_truc_id"],
                "ngay_truc_ca_nguon": selection["ngay_truc_id"],
            },
            format="json",
        )
        self.assertEqual(created.status_code, 201, created.data)
        shift_log = SogiaonhancaHC.objects.get(pk=created.data["id"])
        self.assertEqual(shift_log.lich_truc_nguon_id, schedule.id)
        self.assertFalse(shift_log.nguoi_tao_thuoc_bien_che)
        self.assertIsNotNone(shift_log.dong_bo_bien_che_at)

    def test_create_generate_and_submit_schedule(self):
        response = self.client.post("/api/v1/quanlycatruc/lich-truc/", {"nha_may": self.plant.id, "thang": 8, "nam": 2026, "mau_chu_ky": self.template.id}, format="json")
        self.assertEqual(response.status_code, 201, response.data)
        schedule_id = response.data["id"]
        generated = self.client.post(f"/api/v1/quanlycatruc/lich-truc/{schedule_id}/sinh-lich/", {}, format="json")
        self.assertEqual(generated.status_code, 200, generated.data)
        self.assertEqual(len(generated.data["danh_sach_ngay"]), 31)
        submitted = self.client.post(f"/api/v1/quanlycatruc/lich-truc/{schedule_id}/gui-duyet/", {}, format="json")
        self.assertEqual(submitted.status_code, 200, submitted.data)
        self.assertEqual(submitted.data["trang_thai"], "cho_duyet")

    def test_create_schedule_with_custom_cycle_anchor(self):
        custom_operation_cycle = [dict(item) for item in self.template.chu_ky_van_hanh]
        custom_operation_cycle[0] = {"day": "A", "night": "B", "transition": True}
        custom_admin_cycle = [dict(item) for item in self.template.chu_ky_hanh_chinh]
        response = self.client.post("/api/v1/quanlycatruc/lich-truc/", {
            "nha_may": self.plant.id, "thang": 9, "nam": 2026,
            "mau_chu_ky": self.template.id,
            "che_do_sinh_chu_ky": "moc_tuy_chon",
            "ngay_moc_chu_ky": "2026-09-03",
            "vi_tri_moc_van_hanh": 0,
            "vi_tri_moc_hanh_chinh": 3,
            "chu_ky_van_hanh_tuy_chon": custom_operation_cycle,
            "chu_ky_hanh_chinh_tuy_chon": custom_admin_cycle,
        }, format="json")
        self.assertEqual(response.status_code, 201, response.data)
        self.assertEqual(response.data["che_do_sinh_chu_ky"], "moc_tuy_chon")
        generated = self.client.post(
            f"/api/v1/quanlycatruc/lich-truc/{response.data['id']}/sinh-lich/", {}, format="json"
        )
        self.assertEqual(generated.status_code, 200, generated.data)
        anchor_day = next(item for item in generated.data["danh_sach_ngay"] if item["ngay"] == "2026-09-03")
        self.assertEqual((anchor_day["kip_ca_ngay_ma"], anchor_day["kip_ca_dem_ma"]), ("A", "B"))
        self.assertTrue(anchor_day["la_ngay_chuyen_kip"])

        invalid = self.client.post("/api/v1/quanlycatruc/lich-truc/", {
            "nha_may": self.plant.id, "thang": 10, "nam": 2026,
            "mau_chu_ky": self.template.id,
            "che_do_sinh_chu_ky": "moc_tuy_chon",
        }, format="json")
        self.assertEqual(invalid.status_code, 400, invalid.data)

    def test_create_new_schedule_allocates_next_version(self):
        payload = {
            "nha_may": self.plant.id,
            "thang": 8,
            "nam": 2026,
            "mau_chu_ky": self.template.id,
        }
        first = self.client.post(
            "/api/v1/quanlycatruc/lich-truc/", payload, format="json"
        )
        second = self.client.post(
            "/api/v1/quanlycatruc/lich-truc/", payload, format="json"
        )

        self.assertEqual(first.status_code, 201, first.data)
        self.assertEqual(first.data["phien_ban"], 1)
        self.assertEqual(second.status_code, 201, second.data)
        self.assertEqual(second.data["phien_ban"], 2)

    def test_create_custom_schedule_with_lunar_dates(self):
        source = self.client.post(
            "/api/v1/quanlycatruc/lich-truc/",
            {"nha_may": self.plant.id, "thang": 2, "nam": 2026, "mau_chu_ky": self.template.id},
            format="json",
        )
        self.assertEqual(source.status_code, 201, source.data)
        generated_source = self.client.post(
            f"/api/v1/quanlycatruc/lich-truc/{source.data['id']}/sinh-lich/", {}, format="json"
        )
        self.assertEqual(generated_source.status_code, 200, generated_source.data)
        custom = self.client.post(
            "/api/v1/quanlycatruc/lich-truc/",
            {
                "nha_may": self.plant.id,
                "loai_lich": "chuyen_de",
                "ten_lich": "Lịch trực Tết Nguyên đán 2026",
                "tu_ngay": "2026-02-16",
                "den_ngay": "2026-02-18",
                "lich_thang_goc": source.data["id"],
                "thang": 2,
                "nam": 2026,
                "mau_chu_ky": self.template.id,
            },
            format="json",
        )
        self.assertEqual(custom.status_code, 201, custom.data)
        generated = self.client.post(
            f"/api/v1/quanlycatruc/lich-truc/{custom.data['id']}/sinh-lich/", {}, format="json"
        )
        self.assertEqual(generated.status_code, 200, generated.data)
        self.assertEqual(len(generated.data["danh_sach_ngay"]), 3)
        tet = next(day for day in generated.data["danh_sach_ngay"] if day["ngay"] == "2026-02-17")
        self.assertEqual((tet["ngay_am"], tet["thang_am"]), (1, 1))
        custom_schedule = LichTrucCa.objects.get(pk=custom.data["id"])
        first_day = custom_schedule.danh_sach_ngay.first()
        ThanhVienKipTruc.objects.create(
            kip_truc=first_day.kip_ca_ngay, user=self.user,
            vai_tro="truong_ca", tu_ngay=date(2026, 1, 1),
        )
        custom_schedule.trang_thai = LichTrucCa.TrangThai.DA_DUYET
        custom_schedule.save(update_fields=["trang_thai", "updated_at"])
        excel = self.client.get(f"/api/v1/quanlycatruc/lich-truc/{custom_schedule.id}/xuat-excel/")
        pdf = self.client.get(f"/api/v1/quanlycatruc/lich-truc/{custom_schedule.id}/xuat-pdf/")
        self.assertEqual(excel.status_code, 200)
        self.assertEqual(pdf.status_code, 200)
        from openpyxl import load_workbook
        workbook = load_workbook(BytesIO(excel.content), data_only=True)
        values = [str(cell.value) for row in workbook["Lịch trực ca"].iter_rows() for cell in row if cell.value]
        self.assertNotIn("DANH SÁCH NHÂN SỰ CA TRỰC", values)
        self.assertIn("scheduler", values)

        unit = DonViToChuc.objects.create(
            ma_don_vi="SH-TET", ten_don_vi="Phân công Tết", nha_may=self.plant,
        )
        department = BoPhan.objects.create(
            don_vi=unit, ma_bo_phan="HC-TET", ten_bo_phan="Hành chính Tết",
        )
        holiday_staff = NhanSu.objects.create(
            ho_ten="Nhân sự trực Tết", don_vi=unit, bo_phan=department,
        )
        hc_team = KipTruc.objects.get(nha_may=self.plant, ma_kip="HC1")
        ThanhVienKipTruc.objects.create(
            kip_truc=hc_team, nhan_su=holiday_staff,
            vai_tro="nhan_vien", tu_ngay=date(2026, 1, 1),
        )
        hc2_team = KipTruc.objects.get(nha_may=self.plant, ma_kip="HC2")
        optional_hc_staff = NhanSu.objects.create(
            ho_ten="Nhân sự HC không tham gia", don_vi=unit, bo_phan=department,
        )
        ThanhVienKipTruc.objects.create(
            kip_truc=hc2_team, nhan_su=optional_hc_staff,
            vai_tro="nhan_vien", tu_ngay=date(2026, 1, 1),
        )
        for team in KipTruc.objects.filter(nha_may=self.plant, loai_kip="van_hanh"):
            for index, role in enumerate(("truong_ca", "truc_chinh", "truc_phu"), start=1):
                person = NhanSu.objects.create(
                    ho_ten=f"{team.ma_kip} Tết {index}", don_vi=unit, bo_phan=department,
                )
                ThanhVienKipTruc.objects.create(
                    kip_truc=team, nhan_su=person, vai_tro=role,
                    tu_ngay=date(2026, 1, 1),
                )
        plan = self.client.post(
            "/api/v1/quanlycatruc/phuong-an-phan-cong/",
            {
                "nha_may": self.plant.id,
                "lich_truc": custom_schedule.id,
                "thang": 2,
                "nam": 2026,
                "ngay_hieu_luc": "2026-02-16",
                "ghi_chu": "Phân công riêng cho lịch Tết",
            },
            format="json",
        )
        self.assertEqual(plan.status_code, 201, plan.data)
        self.assertEqual(plan.data["lich_truc"], custom_schedule.id)
        optional_row = next(
            row for row in plan.data["chi_tiet"] if row["nhan_su"] == optional_hc_staff.id
        )
        unassigned = self.client.patch(
            f"/api/v1/quanlycatruc/chi-tiet-phan-cong/{optional_row['id']}/",
            {"kip_truc": None, "vai_tro": ""}, format="json",
        )
        self.assertEqual(unassigned.status_code, 200, unassigned.data)
        applied = self.client.post(
            f"/api/v1/quanlycatruc/phuong-an-phan-cong/{plan.data['id']}/ap-dung/",
            {}, format="json",
        )
        self.assertEqual(applied.status_code, 200, applied.data)
        custom_schedule.refresh_from_db()
        self.assertEqual(
            custom_schedule.snapshot_nhan_su["HC1"][0]["nhan_su_id"],
            holiday_staff.id,
        )
        self.assertEqual(custom_schedule.snapshot_nhan_su["HC2"], [])
        hc2_day = custom_schedule.danh_sach_ngay.filter(kip_hanh_chinh=hc2_team).first()
        self.assertEqual(actual_shift_staff(hc2_day, "hanh_chinh"), [])

        custom_schedule.trang_thai = LichTrucCa.TrangThai.DU_THAO
        custom_schedule.save(update_fields=["trang_thai", "updated_at"])
        customized = self.client.patch(
            f"/api/v1/quanlycatruc/ngay-truc/{hc2_day.id}/",
            {
                "che_do_phan_cong_hc": "tuy_chinh",
                "nhan_su_hc": [{"nhan_su": optional_hc_staff.id, "vai_tro": "nhan_vien"}],
                "ly_do_dieu_chinh": "Bố trí riêng nhân sự KT-HC ngày Tết",
            }, format="json",
        )
        self.assertEqual(customized.status_code, 200, customized.data)
        hc2_day.refresh_from_db()
        self.assertEqual(actual_shift_staff(hc2_day, "hanh_chinh")[0]["nhan_su_id"], optional_hc_staff.id)

        left_empty = self.client.patch(
            f"/api/v1/quanlycatruc/ngay-truc/{hc2_day.id}/",
            {
                "che_do_phan_cong_hc": "khong_bo_tri",
                "nhan_su_hc": [],
                "ly_do_dieu_chinh": "Ngày không bố trí KT-HC",
            }, format="json",
        )
        self.assertEqual(left_empty.status_code, 200, left_empty.data)
        hc2_day.refresh_from_db()
        self.assertEqual(actual_shift_staff(hc2_day, "hanh_chinh"), [])

    def test_monthly_assignment_plan_inherits_validates_and_applies_roster(self):
        unit = DonViToChuc.objects.create(ma_don_vi="SH-PC", ten_don_vi="Phân công", nha_may=self.plant)
        department = BoPhan.objects.create(don_vi=unit, ma_bo_phan="VH-PC", ten_bo_phan="Vận hành")
        roles = ("truong_ca", "truc_chinh", "truc_phu")
        people_by_team = {}
        for team in KipTruc.objects.filter(nha_may=self.plant, loai_kip="van_hanh"):
            people_by_team[team.ma_kip] = []
            for index, role in enumerate(roles, start=1):
                person = NhanSu.objects.create(ho_ten=f"{team.ma_kip} {index}", don_vi=unit, bo_phan=department)
                people_by_team[team.ma_kip].append(person)
                ThanhVienKipTruc.objects.create(kip_truc=team, nhan_su=person, vai_tro=role, tu_ngay=date(2026, 8, 1))
        hc_team = KipTruc.objects.get(nha_may=self.plant, ma_kip="HC1")
        hc_person = NhanSu.objects.create(ho_ten="Nhân sự HC", don_vi=unit, bo_phan=department)
        ThanhVienKipTruc.objects.create(kip_truc=hc_team, nhan_su=hc_person, vai_tro="nhan_vien", tu_ngay=date(2026, 8, 1))

        created = self.client.post("/api/v1/quanlycatruc/phuong-an-phan-cong/", {
            "nha_may": self.plant.id, "thang": 9, "nam": 2026,
            "ngay_hieu_luc": "2026-09-01", "ghi_chu": "Kế thừa tháng 8",
        }, format="json")
        self.assertEqual(created.status_code, 201, created.data)
        self.assertEqual(len(created.data["chi_tiet"]), 13)
        next_version = self.client.post("/api/v1/quanlycatruc/phuong-an-phan-cong/", {
            "nha_may": self.plant.id, "thang": 9, "nam": 2026,
            "ngay_hieu_luc": "2026-09-01", "ghi_chu": "Phương án bổ sung",
        }, format="json")
        self.assertEqual(next_version.status_code, 201, next_version.data)
        self.assertEqual(next_version.data["phien_ban"], 2)
        a_replacement = people_by_team["A"][2]
        a_team = KipTruc.objects.get(nha_may=self.plant, ma_kip="A")
        extra_person = NhanSu.objects.create(ho_ten="Nhân sự tăng cường A", don_vi=unit, bo_phan=department)
        ChiTietPhuongAnPhanCongCa.objects.create(
            phuong_an_id=created.data["id"], nhan_su=extra_person,
            kip_truc=a_team, vai_tro="truc_phu",
        )
        rows = {item["nhan_su"]: item for item in created.data["chi_tiet"]}
        moved_to_hc = self.client.patch(f"/api/v1/quanlycatruc/chi-tiet-phan-cong/{rows[a_replacement.id]['id']}/", {"kip_truc": hc_team.id, "vai_tro": "nhan_vien"}, format="json")
        moved_to_a = self.client.patch(f"/api/v1/quanlycatruc/chi-tiet-phan-cong/{rows[hc_person.id]['id']}/", {"kip_truc": a_team.id, "vai_tro": "truc_phu"}, format="json")
        self.assertEqual(moved_to_hc.status_code, 200, moved_to_hc.data)
        self.assertEqual(moved_to_a.status_code, 200, moved_to_a.data)
        checked = self.client.get(f"/api/v1/quanlycatruc/phuong-an-phan-cong/{created.data['id']}/kiem-tra/")
        self.assertEqual(checked.status_code, 200, checked.data)
        self.assertTrue(checked.data["hop_le"])
        self.assertEqual(len(checked.data["canh_bao"]), 1)
        self.assertIn("Kíp A", checked.data["canh_bao"][0])
        applied = self.client.post(f"/api/v1/quanlycatruc/phuong-an-phan-cong/{created.data['id']}/ap-dung/", {}, format="json")
        self.assertEqual(applied.status_code, 200, applied.data)
        self.assertEqual(applied.data["trang_thai"], "da_ap_dung")
        self.assertTrue(ThanhVienKipTruc.objects.filter(nhan_su=a_replacement, kip_truc=hc_team, tu_ngay=date(2026, 9, 1)).exists())
        self.assertTrue(ThanhVienKipTruc.objects.filter(nhan_su=hc_person, kip_truc=a_team, vai_tro="truc_phu", tu_ngay=date(2026, 9, 1)).exists())

        # An applied plan remains editable until it is explicitly locked.
        moved_back_to_a = self.client.patch(
            f"/api/v1/quanlycatruc/chi-tiet-phan-cong/{rows[a_replacement.id]['id']}/",
            {"kip_truc": a_team.id, "vai_tro": "truc_phu"}, format="json",
        )
        moved_back_to_hc = self.client.patch(
            f"/api/v1/quanlycatruc/chi-tiet-phan-cong/{rows[hc_person.id]['id']}/",
            {"kip_truc": hc_team.id, "vai_tro": "nhan_vien"}, format="json",
        )
        self.assertEqual(moved_back_to_a.status_code, 200, moved_back_to_a.data)
        self.assertEqual(moved_back_to_hc.status_code, 200, moved_back_to_hc.data)
        reapplied = self.client.post(
            f"/api/v1/quanlycatruc/phuong-an-phan-cong/{created.data['id']}/ap-dung/",
            {}, format="json",
        )
        self.assertEqual(reapplied.status_code, 200, reapplied.data)
        self.assertTrue(ThanhVienKipTruc.objects.filter(
            nhan_su=a_replacement, kip_truc=a_team, tu_ngay=date(2026, 9, 1)
        ).exists())
        self.assertTrue(ThanhVienKipTruc.objects.filter(
            nhan_su=hc_person, kip_truc=hc_team, tu_ngay=date(2026, 9, 1)
        ).exists())

        locked = self.client.post(
            f"/api/v1/quanlycatruc/phuong-an-phan-cong/{created.data['id']}/khoa/",
            {}, format="json",
        )
        self.assertEqual(locked.status_code, 200, locked.data)
        rejected = self.client.patch(
            f"/api/v1/quanlycatruc/chi-tiet-phan-cong/{rows[hc_person.id]['id']}/",
            {"kip_truc": a_team.id, "vai_tro": "truc_phu"}, format="json",
        )
    def test_can_delete_draft_schedule_with_attached_assignment_plan(self):
        schedule = self.make_schedule(month=9, year=2026)
        plan = PhuongAnPhanCongCa.objects.create(
            nha_may=self.plant,
            lich_truc=schedule,
            thang=9,
            nam=2026,
            ngay_hieu_luc=date(2026, 9, 1),
            nguoi_tao=self.user,
        )
        response = self.client.delete(f"/api/v1/quanlycatruc/lich-truc/{schedule.id}/")
        self.assertEqual(response.status_code, 204)
        self.assertFalse(LichTrucCa.objects.filter(pk=schedule.id).exists())

    def test_user_cannot_access_another_factory(self):
        schedule = self.make_schedule()
        other_user = create_user("other", self.other_plant, can_view_shift_schedule=True)
        self.client.force_authenticate(other_user)
        response = self.client.get(f"/api/v1/quanlycatruc/lich-truc/{schedule.id}/")
        self.assertEqual(response.status_code, 404)

    def test_unapproved_schedule_cannot_be_exported(self):
        schedule = self.make_schedule()
        response = self.client.get(f"/api/v1/quanlycatruc/lich-truc/{schedule.id}/xuat-excel/")
        self.assertEqual(response.status_code, 400)

    def test_approved_schedule_exports_excel_and_pdf(self):
        schedule = self.make_schedule()
        schedule.ghi_chu = "Ưu tiên bảo đảm vận hành an toàn"
        schedule.save(update_fields=["ghi_chu", "updated_at"])
        ThanhVienKipTruc.objects.create(
            kip_truc=KipTruc.objects.get(nha_may=self.plant, ma_kip="A"),
            user=self.user,
            vai_tro="truong_ca",
            tu_ngay=date(2026, 1, 1),
        )
        hc_user = create_user("hc-export", self.plant)
        ThanhVienKipTruc.objects.create(
            kip_truc=KipTruc.objects.get(nha_may=self.plant, ma_kip="HC1"),
            user=hc_user,
            vai_tro="nhan_vien",
            tu_ngay=date(2026, 1, 1),
        )
        generate_monthly_schedule(schedule.id, self.user)
        transition_schedule(schedule.id, self.user, "gui_duyet")
        transition_schedule(schedule.id, self.user, "phe_duyet")
        excel = self.client.get(f"/api/v1/quanlycatruc/lich-truc/{schedule.id}/xuat-excel/")
        pdf = self.client.get(f"/api/v1/quanlycatruc/lich-truc/{schedule.id}/xuat-pdf/")
        self.assertEqual(excel.status_code, 200)
        self.assertEqual(excel["Content-Type"], "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")
        self.assertTrue(excel.content.startswith(b"PK"))
        from openpyxl import load_workbook

        workbook = load_workbook(BytesIO(excel.content), data_only=True)
        sheet = workbook["Lịch trực ca"]
        excel_values = [
            str(cell.value)
            for row in sheet.iter_rows()
            for cell in row
            if cell.value is not None
        ]
        self.assertIn("DANH SÁCH NHÂN SỰ CA TRỰC", excel_values)
        self.assertIn("scheduler", excel_values)
        self.assertIn("hc-export", excel_values)
        self.assertIn("Ghi chú lịch: Ưu tiên bảo đảm vận hành an toàn", excel_values)
        label_rows = {
            sheet.cell(row=row, column=1).value: row
            for row in range(1, sheet.max_row + 1)
        }
        self.assertIn("08:00-20:00", label_rows)
        self.assertTrue(sheet.cell(row=label_rows["08:00-20:00"], column=2).fill.fgColor.rgb.endswith("E0E7FF"))
        self.assertTrue(sheet.cell(row=label_rows["Thứ"], column=5).fill.fgColor.rgb.endswith("FEF9C3"))
        self.assertEqual(pdf.status_code, 200)
        self.assertEqual(pdf["Content-Type"], "application/pdf")
        self.assertTrue(pdf.content.startswith(b"%PDF"))
        from pypdf import PdfReader

        pdf_text = "\n".join(page.extract_text() or "" for page in PdfReader(BytesIO(pdf.content)).pages)
        self.assertIn("DANH SÁCH NHÂN SỰ CA TRỰC", pdf_text)
        self.assertIn("scheduler", pdf_text)
        self.assertIn("hc-export", pdf_text)
        self.assertIn("Ưu tiên bảo đảm vận hành an toàn", pdf_text)

    def test_staff_substitution_keeps_operation_shift_fully_staffed(self):
        unit = DonViToChuc.objects.create(ma_don_vi="SH-VH", ten_don_vi="Vận hành Sông Hinh", nha_may=self.plant)
        department = BoPhan.objects.create(don_vi=unit, ma_bo_phan="VH", ten_bo_phan="Vận hành")
        people = [
            NhanSu.objects.create(ma_nhan_vien=f"NV{index}", ho_ten=name, don_vi=unit, bo_phan=department)
            for index, name in enumerate(("Trưởng ca A", "Trực chính A", "Trực phụ A", "Nhân sự hành chính"), start=1)
        ]
        team_a = KipTruc.objects.get(nha_may=self.plant, ma_kip="A")
        team_hc = KipTruc.objects.get(nha_may=self.plant, ma_kip="HC1")
        for person, role in zip(people[:3], ("truong_ca", "truc_chinh", "truc_phu")):
            ThanhVienKipTruc.objects.create(kip_truc=team_a, nhan_su=person, vai_tro=role, tu_ngay=date(2026, 1, 1))
        ThanhVienKipTruc.objects.create(kip_truc=team_hc, nhan_su=people[3], vai_tro="nhan_vien", tu_ngay=date(2026, 1, 1))
        schedule = self.make_schedule()
        generate_monthly_schedule(schedule.id, self.user)
        day = schedule.danh_sach_ngay.get(ngay=date(2026, 8, 1))

        created = self.client.post("/api/v1/quanlycatruc/dieu-chinh-nhan-su/", {
            "ngay_truc": day.id,
            "loai_ca": "ca_dem",
            "loai_dieu_chinh": "truc_thay",
            "nhan_su_vang": people[2].id,
            "nhan_su_thay": people[3].id,
            "vai_tro_thay": "truc_phu",
            "ly_do": "Trực phụ nghỉ phép",
        }, format="json")
        self.assertEqual(created.status_code, 201, created.data)
        approved = self.client.post(f"/api/v1/quanlycatruc/dieu-chinh-nhan-su/{created.data['id']}/phe-duyet/", {}, format="json")
        self.assertEqual(approved.status_code, 200, approved.data)
        day_data = self.client.get(f"/api/v1/quanlycatruc/ngay-truc/{day.id}/").data
        actual_names = [item["ho_ten"] for item in day_data["nhan_su_thuc_te"]["ca_dem"]]
        self.assertNotIn("Trực phụ A", actual_names)
        self.assertIn("Nhân sự hành chính", actual_names)
        self.assertFalse(day_data["canh_bao_nhan_su"]["ca_dem"])

    def test_operation_leave_without_replacement_cannot_be_approved(self):
        unit = DonViToChuc.objects.create(ma_don_vi="SH-VH2", ten_don_vi="Vận hành 2", nha_may=self.plant)
        department = BoPhan.objects.create(don_vi=unit, ma_bo_phan="VH2", ten_bo_phan="Vận hành 2")
        roles = ("truong_ca", "truc_chinh", "truc_phu")
        people = []
        team_a = KipTruc.objects.get(nha_may=self.plant, ma_kip="A")
        for index, role in enumerate(roles, start=1):
            person = NhanSu.objects.create(ma_nhan_vien=f"NP{index}", ho_ten=f"Người {index}", don_vi=unit, bo_phan=department)
            people.append(person)
            ThanhVienKipTruc.objects.create(kip_truc=team_a, nhan_su=person, vai_tro=role, tu_ngay=date(2026, 1, 1))
        schedule = self.make_schedule()
        generate_monthly_schedule(schedule.id, self.user)
        day = schedule.danh_sach_ngay.get(ngay=date(2026, 8, 1))
        created = self.client.post("/api/v1/quanlycatruc/dieu-chinh-nhan-su/", {
            "ngay_truc": day.id, "loai_ca": "ca_dem", "loai_dieu_chinh": "nghi_phep",
            "nhan_su_vang": people[2].id, "ly_do": "Nghỉ phép",
        }, format="json")
        self.assertEqual(created.status_code, 201, created.data)
        approved = self.client.post(f"/api/v1/quanlycatruc/dieu-chinh-nhan-su/{created.data['id']}/phe-duyet/", {}, format="json")
        self.assertEqual(approved.status_code, 400)

    def test_two_way_swap_updates_both_operation_shifts(self):
        unit = DonViToChuc.objects.create(ma_don_vi="SH-DOI", ten_don_vi="Vận hành đổi ca", nha_may=self.plant)
        department = BoPhan.objects.create(don_vi=unit, ma_bo_phan="DOI", ten_bo_phan="Vận hành")
        schedule = self.make_schedule()
        generate_monthly_schedule(schedule.id, self.user)
        days = list(schedule.danh_sach_ngay.select_related("kip_ca_ngay").order_by("ngay"))
        source_day = days[0]
        target_day = next(item for item in days[1:] if item.kip_ca_ngay_id != source_day.kip_ca_ngay_id)
        roles = ("truong_ca", "truc_chinh", "truc_phu")
        source_people = []
        target_people = []
        for prefix, team, bucket in (("S", source_day.kip_ca_ngay, source_people), ("T", target_day.kip_ca_ngay, target_people)):
            for index, role in enumerate(roles, start=1):
                person = NhanSu.objects.create(ma_nhan_vien=f"{prefix}{index}", ho_ten=f"{prefix}-{role}", don_vi=unit, bo_phan=department)
                bucket.append(person)
                ThanhVienKipTruc.objects.create(kip_truc=team, nhan_su=person, vai_tro=role, tu_ngay=date(2026, 1, 1))

        created = self.client.post("/api/v1/quanlycatruc/dieu-chinh-nhan-su/", {
            "ngay_truc": source_day.id, "loai_ca": "ca_ngay", "loai_dieu_chinh": "doi_ca",
            "ngay_truc_doi": target_day.id, "loai_ca_doi": "ca_ngay",
            "nhan_su_vang": source_people[2].id, "nhan_su_thay": target_people[2].id,
            "vai_tro_thay": "truc_phu", "vai_tro_doi": "truc_phu", "ly_do": "Hai nhân viên đổi ca",
        }, format="json")
        self.assertEqual(created.status_code, 201, created.data)
        approved = self.client.post(f"/api/v1/quanlycatruc/dieu-chinh-nhan-su/{created.data['id']}/phe-duyet/", {}, format="json")
        self.assertEqual(approved.status_code, 200, approved.data)
        source_data = self.client.get(f"/api/v1/quanlycatruc/ngay-truc/{source_day.id}/").data
        target_data = self.client.get(f"/api/v1/quanlycatruc/ngay-truc/{target_day.id}/").data
        source_names = {item["ho_ten"] for item in source_data["nhan_su_thuc_te"]["ca_ngay"]}
        target_names = {item["ho_ten"] for item in target_data["nhan_su_thuc_te"]["ca_ngay"]}
        self.assertIn(target_people[2].ho_ten, source_names)
        self.assertNotIn(source_people[2].ho_ten, source_names)
        self.assertIn(source_people[2].ho_ten, target_names)
        self.assertNotIn(target_people[2].ho_ten, target_names)
        self.assertFalse(source_data["canh_bao_nhan_su"]["ca_ngay"])
        self.assertFalse(target_data["canh_bao_nhan_su"]["ca_ngay"])

    def test_create_and_approve_adjustment_batch_for_scheduled_shifts_in_six_days(self):
        unit = DonViToChuc.objects.create(ma_don_vi="SH-DOT", ten_don_vi="Vận hành theo đợt", nha_may=self.plant)
        department = BoPhan.objects.create(don_vi=unit, ma_bo_phan="DOT", ten_bo_phan="Vận hành")
        team_a = KipTruc.objects.get(nha_may=self.plant, ma_kip="A")
        team_hc = KipTruc.objects.get(nha_may=self.plant, ma_kip="HC1")
        people = []
        for index, role in enumerate(("truong_ca", "truc_chinh", "truc_phu"), start=1):
            person = NhanSu.objects.create(ma_nhan_vien=f"DOT{index}", ho_ten=f"Đợt {index}", don_vi=unit, bo_phan=department)
            people.append(person)
            ThanhVienKipTruc.objects.create(kip_truc=team_a, nhan_su=person, vai_tro=role, tu_ngay=date(2026, 1, 1))
        replacement = NhanSu.objects.create(ma_nhan_vien="DOT4", ho_ten="Người trực thay đợt", don_vi=unit, bo_phan=department)
        ThanhVienKipTruc.objects.create(kip_truc=team_hc, nhan_su=replacement, vai_tro="nhan_vien", tu_ngay=date(2026, 1, 1))
        schedule = self.make_schedule()
        generate_monthly_schedule(schedule.id, self.user)

        created = self.client.post("/api/v1/quanlycatruc/dieu-chinh-nhan-su/tao-theo-khoang/", {
            "lich_truc": schedule.id, "tu_ngay": "2026-08-01", "den_ngay": "2026-08-06",
            "tu_gio": "12:00", "den_gio": "12:00",
            "loai_ca": "theo_lich", "loai_dieu_chinh": "truc_thay",
            "nhan_su_vang": people[2].id, "nhan_su_thay": replacement.id,
            "ly_do": "Trực thay trong đợt nghỉ phép",
        }, format="json")
        self.assertEqual(created.status_code, 201, created.data)
        self.assertGreater(created.data["so_ca"], 0)
        details = created.data["chi_tiet"]
        self.assertEqual({item["ma_phieu"] for item in details}, {created.data["ma_phieu"]})
        schedule_days = {item.id: item for item in schedule.danh_sach_ngay.all()}
        for item in details:
            item_day = schedule_days[item["ngay_truc"]]
            shift_hour = 20 if item["loai_ca"] == "ca_dem" else 8
            if item["loai_ca"] == "ca_ngay" and item_day.la_ngay_chuyen_kip:
                shift_hour = 12
            shift_start = datetime.combine(item_day.ngay, time(hour=shift_hour))
            self.assertGreaterEqual(shift_start, datetime(2026, 8, 1, 12))
            self.assertLess(shift_start, datetime(2026, 8, 6, 12))

        approved = self.client.post(f"/api/v1/quanlycatruc/dieu-chinh-nhan-su/{details[0]['id']}/phe-duyet-phieu/", {}, format="json")
        self.assertEqual(approved.status_code, 200, approved.data)
        self.assertEqual(approved.data["so_ca"], created.data["so_ca"])
        self.assertFalse(DieuChinhNhanSuCaTruc.objects.filter(ma_phieu=created.data["ma_phieu"]).exclude(trang_thai="da_duyet").exists())
        first_day_id = details[0]["ngay_truc"]
        first_day = self.client.get(f"/api/v1/quanlycatruc/ngay-truc/{first_day_id}/").data
        hc_names = {item["ho_ten"] for item in first_day["nhan_su_thuc_te"]["hanh_chinh"]}
        self.assertNotIn(replacement.ho_ten, hc_names)

        # Deleting batch adjustments removes all items in the batch and reverts staffing
        deleted = self.client.delete(f"/api/v1/quanlycatruc/dieu-chinh-nhan-su/{details[0]['id']}/xoa-phieu/")
        self.assertEqual(deleted.status_code, 200, deleted.data)
        self.assertFalse(DieuChinhNhanSuCaTruc.objects.filter(ma_phieu=created.data["ma_phieu"]).exists())
        first_day_after = self.client.get(f"/api/v1/quanlycatruc/ngay-truc/{first_day_id}/").data
        op_names_after = {item["ho_ten"] for item in first_day_after["nhan_su_thuc_te"][details[0]["loai_ca"]]}
        self.assertIn(people[2].ho_ten, op_names_after)
        self.assertNotIn(replacement.ho_ten, op_names_after)

    def test_admin_shift_leave_keeps_only_actual_remaining_staff(self):
        unit = DonViToChuc.objects.create(
            ma_don_vi="SH-HC-GIAM", ten_don_vi="Hành chính giảm nhân sự", nha_may=self.plant
        )
        department = BoPhan.objects.create(
            don_vi=unit, ma_bo_phan="HC-GIAM", ten_bo_phan="Hành chính"
        )
        schedule = self.make_schedule()
        generate_monthly_schedule(schedule.id, self.user)
        first_day = schedule.danh_sach_ngay.select_related("kip_hanh_chinh").order_by("ngay").first()
        admin_team = first_day.kip_hanh_chinh
        people = []
        for index, role in enumerate(("truong_ca", "truc_chinh"), start=1):
            person = NhanSu.objects.create(
                ma_nhan_vien=f"HCG{index}", ho_ten=f"Nhân sự HC {index}", don_vi=unit, bo_phan=department
            )
            people.append(person)
            ThanhVienKipTruc.objects.create(
                kip_truc=admin_team, nhan_su=person, vai_tro=role, tu_ngay=date(2026, 1, 1)
            )

        leave = self.client.post("/api/v1/quanlycatruc/dieu-chinh-nhan-su/tao-theo-khoang/", {
            "lich_truc": schedule.id,
            "tu_ngay": first_day.ngay.isoformat(), "den_ngay": first_day.ngay.isoformat(),
            "tu_gio": "08:00", "den_gio": "20:00",
            "loai_ca": "hanh_chinh", "loai_dieu_chinh": "nghi_bu",
            "nhan_su_vang": people[0].id,
            "ly_do": "Nghỉ bù theo kế hoạch",
        }, format="json")
        self.assertEqual(leave.status_code, 201, leave.data)
        self.assertEqual(leave.data["so_ca"], 1)
        approved = self.client.post(
            f"/api/v1/quanlycatruc/dieu-chinh-nhan-su/{leave.data['chi_tiet'][0]['id']}/phe-duyet-phieu/",
            {}, format="json",
        )
        self.assertEqual(approved.status_code, 200, approved.data)
        remaining_ids = {item["nhan_su_id"] for item in actual_shift_staff(first_day, "hanh_chinh")}
        self.assertEqual(remaining_ids, {people[1].id})

    def test_adjustment_batch_rejects_more_than_six_days(self):
        schedule = self.make_schedule()
        response = self.client.post("/api/v1/quanlycatruc/dieu-chinh-nhan-su/tao-theo-khoang/", {
            "lich_truc": schedule.id, "tu_ngay": "2026-08-01", "den_ngay": "2026-08-07",
            "loai_ca": "theo_lich", "loai_dieu_chinh": "nghi_phep",
            "nhan_su_vang": 999999, "ly_do": "Nghỉ dài ngày",
        }, format="json")
        self.assertEqual(response.status_code, 400)

    def test_schedule_list_filters_by_month_and_year(self):
        august = self.make_schedule()
        self.make_schedule(month=9)
        response = self.client.get(f"/api/v1/quanlycatruc/lich-truc/?nha_may={self.plant.id}&thang=8&nam=2026")
        self.assertEqual(response.status_code, 200)
        items = response.data.get("results", response.data) if isinstance(response.data, dict) else response.data
        self.assertEqual([item["id"] for item in items], [august.id])

    def test_draft_schedule_can_be_updated_and_deleted(self):
        schedule = self.make_schedule()
        updated = self.client.patch(
            f"/api/v1/quanlycatruc/lich-truc/{schedule.id}/", {"ghi_chu": "Điều chỉnh nhân sự"}, format="json"
        )
        self.assertEqual(updated.status_code, 200, updated.data)
        self.assertEqual(updated.data["ghi_chu"], "Điều chỉnh nhân sự")
        deleted = self.client.delete(f"/api/v1/quanlycatruc/lich-truc/{schedule.id}/")
        self.assertEqual(deleted.status_code, 204)
        self.assertFalse(LichTrucCa.objects.filter(pk=schedule.id).exists())

    def test_user_cannot_create_member_for_another_factory(self):
        other_team = KipTruc.objects.create(
            nha_may=self.other_plant, loai_kip="van_hanh", ma_kip="A", ten_kip="Kíp A", thu_tu=1
        )
        response = self.client.post(
            "/api/v1/quanlycatruc/thanh-vien-kip/",
            {"kip_truc": other_team.id, "user": self.user.id, "vai_tro": "nhan_vien", "tu_ngay": "2026-08-01"},
            format="json",
        )
        self.assertEqual(response.status_code, 403)

    def test_non_staff_roster_manager_can_list_personnel_without_accounts(self):
        unit = DonViToChuc.objects.create(ma_don_vi="SH", ten_don_vi="Nhà máy Sông Hinh", nha_may=self.plant)
        department = BoPhan.objects.create(don_vi=unit, ma_bo_phan="VH", ten_bo_phan="Vận hành", loai_bo_phan="van_hanh")
        personnel = NhanSu.objects.create(ho_ten="Nguyễn Văn Trực", don_vi=unit, bo_phan=department, tu_ngay=date(2026, 1, 1))
        self.assertFalse(self.user.is_staff)

        response = self.client.get(f"/api/v1/quanlycatruc/kip-truc/nhan-su-options/?nha_may={self.plant.id}")

        self.assertEqual(response.status_code, 200, response.data)
        returned_ids = {item["id"] for item in response.data}
        self.assertIn(personnel.id, returned_ids)
        self.assertIsNone(response.data[0]["user"])

    def test_can_create_personnel_without_user_and_assign_to_team(self):
        unit_response = self.client.post("/api/v1/quanlycatruc/don-vi/", {
            "ma_don_vi": "SH", "ten_don_vi": "Nhà máy Sông Hinh", "loai_don_vi": "nha_may", "nha_may": self.plant.id
        }, format="json")
        self.assertEqual(unit_response.status_code, 201, unit_response.data)
        department_response = self.client.post("/api/v1/quanlycatruc/bo-phan/", {
            "don_vi": unit_response.data["id"], "ma_bo_phan": "BV", "ten_bo_phan": "Bảo vệ", "loai_bo_phan": "bao_ve"
        }, format="json")
        self.assertEqual(department_response.status_code, 201, department_response.data)
        personnel_response = self.client.post("/api/v1/quanlycatruc/nhan-su/", {
            "ho_ten": "Trần Văn An", "don_vi": unit_response.data["id"], "bo_phan": department_response.data["id"], "tu_ngay": "2026-01-01"
        }, format="json")
        self.assertEqual(personnel_response.status_code, 201, personnel_response.data)
        self.assertIsNone(personnel_response.data["user"])

        team = KipTruc.objects.get(nha_may=self.plant, ma_kip="A")
        member_response = self.client.post("/api/v1/quanlycatruc/thanh-vien-kip/", {
            "kip_truc": team.id, "nhan_su": personnel_response.data["id"], "vai_tro": "nhan_vien", "tu_ngay": "2026-08-01"
        }, format="json")
        self.assertEqual(member_response.status_code, 201, member_response.data)
        membership = ThanhVienKipTruc.objects.get(pk=member_response.data["id"])
        self.assertEqual(membership.nhan_su.ho_ten, "Trần Văn An")
        self.assertIsNone(membership.user_id)
